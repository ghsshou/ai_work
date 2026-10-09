#!/usr/bin/env python3
"""Build the design-time AI cluster simulator workload-engine slides."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = "docs/design-simulator/设计态仿真器_负载引擎立项.pptx"
FONT = "微软雅黑"

BG = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
DIM = RGBColor(0x7A, 0x88, 0x96)
LINE = RGBColor(0xD0, 0xD8, 0xE0)
BLUE = RGBColor(0x1A, 0x5F, 0xA8)
BLUE_DK = RGBColor(0x12, 0x3E, 0x6E)
BLUE_BG = RGBColor(0xE3, 0xEE, 0xF8)
CORE = RGBColor(0xC4, 0x5A, 0x10)
CORE_BG = RGBColor(0xFD, 0xF3, 0xEA)
CORE_LT = RGBColor(0xFB, 0xE3, 0xCF)
GREEN = RGBColor(0x3D, 0x86, 0x2C)
GREEN_BG = RGBColor(0xE8, 0xF3, 0xE2)
GRAY_BG = RGBColor(0xE9, 0xEC, 0xF0)
HEAD = RGBColor(0x1F, 0x3A, 0x5C)
ROW_ALT = RGBColor(0xF7, 0xF9, 0xFB)

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 2


def _ea_font(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", name)


def _style(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    run.font.italic = False
    _ea_font(run)


def _anchor(tf, anchor):
    tf._txBody.bodyPr.set("anchor", {
        MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b",
    }.get(anchor, "t"))


def _margins(tf, lr=0.05, tb=0.03):
    tf.margin_left = tf.margin_right = Inches(lr)
    tf.margin_top = tf.margin_bottom = Inches(tb)


def text(slide, l, t, w, h, content, size=11, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """content: str, or list of runs-lines.

    A line is either a str or a list of (text, size, bold, color) segments.
    """
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    _margins(tf)
    _anchor(tf, anchor)
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2.5)
        p.space_before = Pt(0)
        segs = line if isinstance(line, list) else [(line, size, bold, color)]
        for seg in segs:
            run = p.add_run()
            run.text = seg[0]
            _style(run, seg[1], seg[2], seg[3])
    return box


def shape(slide, kind, l, t, w, h, fill, line=None, line_pt=0.9, adj=None):
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_pt)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.shadow.inherit = False
    return sh


def rect(slide, l, t, w, h, fill, line=None, line_pt=0.9):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line, line_pt)


def card(slide, l, t, w, h, fill=WHITE, line=LINE, line_pt=0.9, adj=0.05):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line, line_pt, adj)


def label_shape(sh, content, size, bold, color, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    _margins(tf, 0.03, 0.01)
    _anchor(tf, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = content
    _style(run, size, bold, color)


def arrow(slide, l, t, w, h, color=BLUE, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, l, t, w, h, color, adj=0.5)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, BG)
    return slide


def footer(slide, page, note):
    rect(slide, 0, Inches(7.26), W, Inches(0.24), GRAY_BG)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), BLUE)
    text(slide, Inches(0.36), Inches(7.26), Inches(11), Inches(0.24), note,
         8.5, False, DIM, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(11.3), Inches(7.26), Inches(1.7), Inches(0.24),
         f"{page:02d}  /  {TOTAL:02d}", 8.5, False, MUTED, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def title_block(slide, kicker, title, subtitle):
    rect(slide, 0, 0, W, Inches(0.07), BLUE)
    text(slide, Inches(0.36), Inches(0.13), Inches(12.6), Inches(0.26), kicker, 10.5, True, BLUE)
    text(slide, Inches(0.36), Inches(0.36), Inches(12.6), Inches(0.46), title, 21, True, INK)
    text(slide, Inches(0.36), Inches(0.82), Inches(12.6), Inches(0.32), subtitle, 11.5, False, MUTED)


def bullets(slide, l, t, w, h, items, size=9.5, color=INK, dot=CORE):
    lines = []
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            lines.append([("▪ ", size, True, dot), (head, size, True, INK), (body, size, False, color)])
        else:
            lines.append([("▪ ", size, True, dot), (it, size, False, color)])
    return text(slide, l, t, w, h, lines, size)


def module_card(slide, l, t, w, h, num, title, en, items):
    card(slide, l, t, w, h, WHITE, CORE, 1.0)
    head = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, Inches(0.36), CORE_LT, adj=0.2)
    text(slide, l + Inches(0.06), t + Inches(0.02), w - Inches(0.1), Inches(0.32),
         [[(num + " ", 11.5, True, CORE), (title, 11.5, True, INK), ("  " + en if en else "", 8, False, MUTED)]],
         anchor=MSO_ANCHOR.MIDDLE)
    bullets(slide, l + Inches(0.06), t + Inches(0.42), w - Inches(0.1), h - Inches(0.46), items)
    return head


def slide_main(prs):
    s = blank(prs)
    title_block(
        s, "设计态 AI 集群仿真器  ·  立项核心",
        "以「负载引擎」为核心的设计态仿真器：负载刻画准，系统设计才可信",
        "底层 NPU / 内存存储 / 网络仿真能力已具备；本项目聚焦负载的画像、到达建模、Trace 拟合外推、数据工厂与统一导入，"
        "形成「真实负载 → 可外推负载 → 设计结论」闭环。")

    top = Inches(1.25)
    bottom = Inches(6.42)

    # ---- Left: inputs + design-time traits ----
    lx, lw = Inches(0.36), Inches(2.18)
    text(s, lx, top, lw, Inches(0.28), "输入源", 11.5, True, BLUE_DK)
    inputs = [
        ("生产 Trace", "训练 Profiling（msprof/Kineto）、HCCL 通信日志、推理服务请求日志、集群遥测"),
        ("模型与并行配置", "模型结构定义、并行策略、目标模型假设（如 10T MoE、百万上下文）"),
        ("业务规划", "流量预测、SLO 等级、租户结构、训推混部比例"),
    ]
    y = top + Inches(0.32)
    for head, body in inputs:
        card(s, lx, y, lw, Inches(0.92), BLUE_BG, None)
        text(s, lx + Inches(0.05), y + Inches(0.04), lw - Inches(0.08), Inches(0.86),
             [[(head, 10, True, BLUE_DK)], [(body, 8.8, False, INK)]])
        y += Inches(1.0)

    ty = y + Inches(0.04)
    th = bottom - ty
    card(s, lx, ty, lw, th, WHITE, BLUE, 1.0)
    text(s, lx + Inches(0.06), ty + Inches(0.04), lw - Inches(0.1), th - Inches(0.06), [
        [("设计态 ≠ 在线工具", 10, True, BLUE_DK)],
        [("▪ ", 8.8, True, BLUE), ("无实时约束：以时间换精度，可做包级细粒度仿真", 8.8, False, INK)],
        [("▪ ", 8.8, True, BLUE), ("目标系统与模型往往尚不存在：负载必须", 8.8, False, INK),
         ("可参数化、可外推", 8.8, True, CORE), ("，不能只靠回放", 8.8, False, INK)],
        [("▪ ", 8.8, True, BLUE), ("服务于设计空间探索（DSE）：需批量扫描上千种配置", 8.8, False, INK)],
    ])

    arrow(s, Inches(2.58), Inches(3.55), Inches(0.24), Inches(0.42), CORE)

    # ---- Center: workload engine ----
    cx, cw = Inches(2.86), Inches(7.12)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx, top, cw, bottom - top, CORE_BG, CORE, 1.75, 0.025)
    text(s, cx + Inches(0.1), top + Inches(0.03), cw - Inches(0.2), Inches(0.32),
         [[("负载引擎 Workload Engine", 13, True, CORE),
           ("   本项目核心：让负载可度量、可生成、可校准", 9.5, False, MUTED)]],
         anchor=MSO_ANCHOR.MIDDLE)

    gap = Inches(0.08)
    mx = cx + Inches(0.1)
    mw = (cw - Inches(0.2) - 2 * gap) / 3
    my, mh = top + Inches(0.4), Inches(2.0)
    module_card(s, mx, my, mw, mh, "①", "模型画像", "", [
        ("结构画像：", "Dense/MoE/长序列；层数、专家数、TopK"),
        ("并行展开：", "TP/PP/DP/EP/CP → 每 rank 算子图 + 通信"),
        ("算子代价：", "Roofline + 实测表 + 回归外推"),
        ("动态特征：", "MoE 路由偏斜、KV 增长、重计算"),
    ])
    module_card(s, mx + mw + gap, my, mw, mh, "②", "到达与请求分布", "", [
        ("到达过程：", "突发 MMPP/Hawkes、日周期、多租户叠加"),
        ("请求形态：", "输入/输出长度联合分布、前缀复用、多轮/Agent"),
        ("训练作业：", "到达、规模幂律、时长、故障"),
        ("条件化：", "按业务/时段/SLO 参数化"),
    ])
    module_card(s, mx + 2 * (mw + gap), my, mw, mh, "③", "Trace 拟合与外推", "", [
        ("对齐提取：", "跨 rank 时钟对齐、依赖与关键路径重建"),
        ("压缩：", "代表迭代 + 扰动统计"),
        ("拟合：", "混合分布（EM）+ KS/W1 检验"),
        ("外推：", "千卡→十万卡扩展；跨硬件重定向"),
    ])

    # Data factory pipeline
    fy, fh = my + mh + Inches(0.08), Inches(1.2)
    fx, fw = mx, cw - Inches(0.2)
    card(s, fx, fy, fw, fh, WHITE, CORE, 1.0)
    text(s, fx + Inches(0.06), fy + Inches(0.03), fw, Inches(0.3),
         [[("④ ", 11.5, True, CORE), ("负载数据工厂", 11.5, True, INK), ("  Workload Data Factory", 8, False, MUTED)]],
         anchor=MSO_ANCHOR.MIDDLE)
    stages = ["采集", "清洗/脱敏", "特征化", "统计/生成", "合成", "保真校验", "版本入库"]
    px = fx + Inches(0.08)
    pw = (fw - Inches(0.16)) / len(stages)
    for i, st in enumerate(stages):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = shape(s, kind, px + i * pw, fy + Inches(0.36), pw + Inches(0.04), Inches(0.34),
                   CORE if i in (3, 5) else CORE_LT, adj=0.3)
        label_shape(sh, st, 8.3, True, WHITE if i in (3, 5) else INK)
    text(s, fx + Inches(0.06), fy + Inches(0.74), fw - Inches(0.1), Inches(0.46), [
        [("产出：", 8.8, True, CORE),
         ("场景化标准负载库：预训练 / RL 后训练 / 在线推理 / PD 分离 / Agent 长上下文 / 训推混部", 8.8, False, INK)],
        [("手段：", 8.8, True, CORE),
         ("参数化生成器 + 生成式合成，按「模型×规模×流量」旋钮批量产出 DSE 输入", 8.8, False, INK)],
    ])

    # IR + calibration
    iy = fy + fh + Inches(0.08)
    ih = bottom - Inches(0.08) - iy
    iw = Inches(4.42)
    card(s, fx, iy, iw, ih, WHITE, CORE, 1.0)
    text(s, fx + Inches(0.06), iy + Inches(0.03), iw - Inches(0.1), ih - Inches(0.05), [
        [("⑤ ", 11.5, True, CORE), ("统一负载 IR 与多保真导入", 11.5, True, INK), ("  Workload IR", 8, False, MUTED)],
        [("▪ ", 8.8, True, CORE), ("硬件无关执行图 IR（兼容 Chakra ET）：计算/通信/访存节点+依赖", 8.8, False, INK)],
        [("▪ ", 8.8, True, CORE), ("L0 解析 → L1 算子图 → L2 包级，热点按需提升保真度", 8.8, False, INK)],
        [("▪ ", 8.8, True, CORE), ("集合通信分解为流下发网络仿真；访存/KV 流量下发内存仿真", 8.8, False, INK)],
    ])
    kx = fx + iw + Inches(0.08)
    kw = fw - iw - Inches(0.08)
    card(s, kx, iy, kw, ih, WHITE, CORE, 1.0)
    text(s, kx + Inches(0.06), iy + Inches(0.03), kw - Inches(0.1), ih - Inches(0.05), [
        [("⑥ ", 11.5, True, CORE), ("闭环校准", 11.5, True, INK), ("  Calibration", 8, False, MUTED)],
        [("▪ ", 8.8, True, CORE), ("小规模实测对标，误差回灌代价模型", 8.8, False, INK)],
        [("▪ ", 8.8, True, CORE), ("双保真：", 8.8, True, INK), ("分布保真 + 效用保真（结论一致）", 8.8, False, INK)],
    ])

    arrow(s, Inches(10.02), Inches(3.55), Inches(0.24), Inches(0.42), CORE)

    # ---- Right: kernels, outputs, benchmarks ----
    rx, rw = Inches(10.3), Inches(2.67)
    text(s, rx, top, rw, Inches(0.28),
         [[("底层仿真内核", 11.5, True, BLUE_DK), ("  已具备", 9, True, GREEN)]])
    ky = top + Inches(0.3)
    kh = Inches(0.42)
    kgap = Inches(0.06)
    kws = (rw - 2 * kgap) / 3
    for i, name in enumerate(["NPU 仿真", "内存/存储", "网络仿真"]):
        sh = card(s, rx + i * (kws + kgap), ky, kws, kh, GRAY_BG, LINE, 0.75, 0.15)
        label_shape(sh, name, 8.8, True, MUTED)

    oy = ky + kh + Inches(0.12)
    text(s, rx, oy, rw, Inches(0.28), "设计输出（DSE）", 11.5, True, BLUE_DK)
    oh = Inches(1.5)
    card(s, rx, oy + Inches(0.3), rw, oh, GREEN_BG, None)
    bullets(s, rx + Inches(0.06), oy + Inches(0.34), rw - Inches(0.1), oh - Inches(0.06), [
        "超节点规模与组网拓扑",
        "Scale-up/Scale-out 带宽配比",
        "HBM/池化内存/存储层级容量",
        "PD 配比、调度与并行策略",
        "瓶颈定位与 TCO/能效评估",
    ], 9, dot=GREEN)

    by = oy + Inches(0.3) + oh + Inches(0.12)
    text(s, rx, by, rw, Inches(0.28), "业界对标与切入点", 11.5, True, BLUE_DK)
    bh = bottom - by - Inches(0.3)
    card(s, rx, by + Inches(0.3), rw, bh, WHITE, LINE)
    text(s, rx + Inches(0.06), by + Inches(0.33), rw - Inches(0.1), bh - Inches(0.05), [
        [("ASTRA-sim 2.0+Chakra", 8.5, True, BLUE), ("：图式执行 Trace IR", 8.5, False, INK)],
        [("SimAI+AICB", 8.5, True, BLUE), ("（阿里）：训练通信负载生成", 8.5, False, INK)],
        [("Vidur", 8.5, True, BLUE), ("（微软）：算子画像+请求分布驱动推理", 8.5, False, INK)],
        [("NVIDIA DSX", 8.5, True, BLUE), ("：AI 工厂数字孪生蓝图", 8.5, False, INK)],
        [("短板：", 8.5, True, CORE),
         ("负载多为回放或简单分布，缺少面向未来模型与超大规模的可外推负载工程", 8.5, False, INK)],
    ])

    # ---- Bottom KPI strip ----
    sy, sh_ = Inches(6.5), Inches(0.68)
    rect(s, Inches(0.36), sy, Inches(12.61), sh_, WHITE, LINE, 0.75)
    rect(s, Inches(0.36), sy, Inches(1.3), sh_, BLUE)
    text(s, Inches(0.36), sy, Inches(1.3), sh_,
         [[("立项目标", 11, True, WHITE)], [("建议值，待对齐", 8, False, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    kpis = [
        ("≤10%", "迭代时间/吞吐仿真误差（千卡实测对标）"),
        ("1K→100K", "千卡 Trace 外推生成十万卡负载"),
        ("6类·30+", "场景类别 · 版本化标准负载"),
        ("≥1000", "配置/轮批量 DSE；解析级分钟、包级小时"),
    ]
    kx0 = Inches(1.74)
    kwid = (Inches(12.97) - kx0) / len(kpis)
    for i, (num, desc) in enumerate(kpis):
        x = kx0 + i * kwid
        if i:
            rect(s, x - Inches(0.02), sy + Inches(0.12), Inches(0.012), sh_ - Inches(0.24), LINE)
        text(s, x, sy, Inches(1.32), sh_, num, 13, True, CORE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(1.3), sy, kwid - Inches(1.34), sh_, desc, 8.8, False, INK,
             anchor=MSO_ANCHOR.MIDDLE)

    footer(s, 1, "设计态 AI 集群仿真器  ·  负载引擎立项  ·  内部讨论稿")
    s.notes_slide.notes_text_frame.text = (
        "讲解主线：底层算力/内存/网络仿真解决的是「系统怎么跑」，而设计态真正的输入是「跑什么」。"
        "设计态的目标系统和未来模型往往还不存在，所以负载不能只回放历史 Trace，必须可画像、可参数化、可外推。\n"
        "①模型画像解决单作业内部的计算/通信/访存结构；②到达与请求分布解决多作业、多请求在时间上的叠加；"
        "③Trace 拟合与外推把真实生产数据变成可放大、可迁移的模型参数；④数据工厂把以上能力流水线化、资产化；"
        "⑤统一 IR 让同一份负载以不同保真度驱动已有仿真内核；⑥闭环校准用实测保证可信度。\n"
        "KPI 为建议值，需结合现有实测平台规模与项目周期对齐。")
    return s


def slide_detail(prs):
    s = blank(prs)
    title_block(
        s, "备份页  ·  核心技术展开",
        "负载引擎六项核心技术：关键问题、技术路线、业界参考与难点",
        "按「单作业结构 → 多作业时间叠加 → 真实数据拟合 → 流水线资产化 → 导入仿真 → 闭环校准」组织。")

    data = [
        ["核心技术", "要回答的关键问题", "技术路线", "业界参考", "难点 / 风险"],
        ["① 模型画像",
         "一个作业在给定并行策略下，每个 rank 何时算什么、传多少、访存多少？",
         "计算图抽取（框架图 dump/ET）→ 并行策略符号化展开 → 算子代价模型（Roofline+实测表+回归）→ MoE 路由/KV 动态注入",
         "Chakra ET、Calculon、Vidur 算子 Profiling、Lumos",
         "未见 shape 与新算子的代价外推；MoE 路由偏斜的刻画"],
        ["② 到达与请求分布",
         "多请求/多作业在时间上如何叠加？突发、周期、长尾如何刻画？",
         "MMPP/Hawkes 刻画突发；输入/输出长度 Copula 联合建模；会话/前缀复用建模；训练作业规模幂律与故障模型",
         "Azure LLM Trace、BurstGPT、Mooncake Trace、ServeGen",
         "生产数据可得性；Agent/RL 等新形态负载缺少历史样本"],
        ["③ Trace 拟合与外推",
         "如何把小规模、旧硬件上的真实 Trace 变成大规模、新硬件的负载？",
         "多源对齐与依赖重建 → 迭代去重压缩 → 分布拟合 + 检验 → rank 对称扩展、跨硬件重定向、模型 what-if 改写",
         "Chakra trace link/converter、SimAI、Echo",
         "外推非线性（拥塞、尾延迟）；跨硬件重定向时依赖关系变化"],
        ["④ 负载数据工厂",
         "如何让负载可复用、可追溯、可批量产出？",
         "采集→清洗脱敏→特征化→建模→合成→校验→入库流水线；参数化生成器 + 生成式合成；负载版本与血缘管理",
         "Mystique（Meta）、MLPerf 负载集",
         "脱敏后的保真度；生成式模型的可解释性"],
        ["⑤ 统一负载 IR 与导入",
         "同一份负载如何以不同保真度驱动 NPU/内存/网络内核？",
         "硬件无关执行图 IR（兼容 Chakra ET）；L0/L1/L2 多保真切换；集合通信分解为流；代表迭代采样加速",
         "ASTRA-sim 2.0、SimAI（ns-3 后端）",
         "多内核协同的时间同步；万卡级图规模与仿真速度"],
        ["⑥ 闭环校准",
         "仿真结论有多可信？误差来自哪里？",
         "小规模实测对标 → 分层误差归因（算子/通信/调度）→ 系数回灌；分布保真 + 效用保真双指标",
         "ASTRA-sim/SimAI 实测验证方法",
         "大规模实测资源有限；误差随规模放大"],
    ]
    l, t = Inches(0.36), Inches(1.25)
    col_w = [Inches(1.7), Inches(2.5), Inches(4.05), Inches(2.2), Inches(2.16)]
    ts = s.shapes.add_table(len(data), len(data[0]), l, t, sum(col_w, Inches(0)), Inches(5.6))
    table = ts.table
    for i, cw in enumerate(col_w):
        table.columns[i].width = cw
    table.rows[0].height = Inches(0.36)
    for r in range(1, len(data)):
        table.rows[r].height = Inches(0.86)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = HEAD
            elif c == 0:
                cell.fill.fore_color.rgb = CORE_LT
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else ROW_ALT
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = val
            if r == 0:
                _style(run, 10.5, True, WHITE)
            elif c == 0:
                _style(run, 10.5, True, CORE)
            else:
                _style(run, 9.2, False, INK)

    footer(s, 2, "参考：ASTRA-sim 2.0/Chakra、SimAI（NSDI'25）、Vidur（MLSys'24）、NVIDIA Omniverse DSX、"
                 "Mystique（ISCA'23）、ServeGen、BurstGPT、Mooncake")
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_main(prs)
    slide_detail(prs)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
