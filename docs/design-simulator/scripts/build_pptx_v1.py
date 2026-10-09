#!/usr/bin/env python3
"""Build the design-time AI cluster simulator workload-engine slides."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = "docs/design-simulator/设计态仿真器_负载引擎立项_v1.pptx"
FONT = "微软雅黑"

BG = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
DIM = RGBColor(0x7A, 0x88, 0x96)
LINE = RGBColor(0xD0, 0xD8, 0xE0)
BLUE = RGBColor(0x1A, 0x5F, 0xA8)
BLUE_DK = RGBColor(0x12, 0x3E, 0x6E)
BLUE_BG = RGBColor(0xE3, 0xEE, 0xF8)
CORE = RGBColor(0xC4, 0x5A, 0x10)
CORE_BG = RGBColor(0xFD, 0xF3, 0xEA)
CORE_LT = RGBColor(0xFB, 0xE3, 0xCF)
GREEN = RGBColor(0x3D, 0x86, 0x2C)
GREEN_BG = RGBColor(0xE8, 0xF3, 0xE2)
GRAY_BG = RGBColor(0xE9, 0xEC, 0xF0)
HEAD = RGBColor(0x1F, 0x3A, 0x5C)
ROW_ALT = RGBColor(0xF7, 0xF9, 0xFB)

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 2


def _ea_font(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", name)


def _style(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    run.font.italic = False
    _ea_font(run)


def _anchor(tf, anchor):
    tf._txBody.bodyPr.set("anchor", {
        MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b",
    }.get(anchor, "t"))


def _margins(tf, lr=0.05, tb=0.03):
    tf.margin_left = tf.margin_right = Inches(lr)
    tf.margin_top = tf.margin_bottom = Inches(tb)


def text(slide, l, t, w, h, content, size=11, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """content: str, or list of runs-lines.

    A line is either a str or a list of (text, size, bold, color) segments.
    """
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    _margins(tf)
    _anchor(tf, anchor)
    lines = content if isinstance(content, list) else content.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2.5)
        p.space_before = Pt(0)
        segs = line if isinstance(line, list) else [(line, size, bold, color)]
        for seg in segs:
            run = p.add_run()
            run.text = seg[0]
            _style(run, seg[1], seg[2], seg[3])
    return box


def shape(slide, kind, l, t, w, h, fill, line=None, line_pt=0.9, adj=None):
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_pt)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.shadow.inherit = False
    return sh


def rect(slide, l, t, w, h, fill, line=None, line_pt=0.9):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line, line_pt)


def card(slide, l, t, w, h, fill=WHITE, line=LINE, line_pt=0.9, adj=0.05):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line, line_pt, adj)


def label_shape(sh, content, size, bold, color, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    _margins(tf, 0.03, 0.01)
    _anchor(tf, MSO_ANCHOR.MIDDLE)
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _style(run, size, bold, color)


def arrow(slide, l, t, w, h, color=BLUE, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, l, t, w, h, color, adj=0.5)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, BG)
    return slide


def footer(slide, page, note):
    rect(slide, 0, Inches(7.26), W, Inches(0.24), GRAY_BG)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), BLUE)
    text(slide, Inches(0.36), Inches(7.26), Inches(11), Inches(0.24), note,
         8.5, False, DIM, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(11.3), Inches(7.26), Inches(1.7), Inches(0.24),
         f"{page:02d}  /  {TOTAL:02d}", 8.5, False, MUTED, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def title_block(slide, kicker, title, subtitle):
    rect(slide, 0, 0, W, Inches(0.07), BLUE)
    text(slide, Inches(0.36), Inches(0.13), Inches(12.6), Inches(0.26), kicker, 10.5, True, BLUE)
    text(slide, Inches(0.36), Inches(0.36), Inches(12.6), Inches(0.46), title, 21, True, INK)
    text(slide, Inches(0.36), Inches(0.82), Inches(12.6), Inches(0.32), subtitle, 11.5, False, MUTED)


def box(slide, l, t, w, h, label, fill, line=None, size=9, bold=False, color=INK, adj=0.12):
    sh = card(slide, l, t, w, h, fill, line, 0.75, adj)
    label_shape(sh, label, size, bold, color)
    return sh


def grid(slide, l, t, w, h, labels, cols, fill, line, size=9, gap=0.06):
    rows = -(-len(labels) // cols)
    g = Inches(gap)
    cw = (w - (cols - 1) * g) / cols
    rh = (h - (rows - 1) * g) / rows
    for i, lab in enumerate(labels):
        r, c = divmod(i, cols)
        box(slide, l + c * (cw + g), t + r * (rh + g), cw, rh, lab, fill, line, size)


def layer_tag(slide, t, h, title, sub, fill):
    sh = rect(slide, Inches(0.36), t, Inches(0.98), h, fill)
    text(slide, Inches(0.36), t, Inches(0.98), h,
         [[(title, 10.5, True, WHITE)]] + ([[(sub, 8, False, WHITE)]] if sub else []),
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return sh


def down_arrow(slide, x, t, h=Inches(0.14), color=DIM):
    return shape(slide, MSO_SHAPE.DOWN_ARROW, x - Inches(0.16), t, Inches(0.32), h, color, adj=0.5)


def module(slide, l, t, w, h, title, labels, cols, title_h=Inches(0.3)):
    card(slide, l, t, w, h, WHITE, CORE, 1.0, 0.04)
    text(slide, l + Inches(0.05), t + Inches(0.02), w - Inches(0.1), title_h, title,
         10, True, CORE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pad = Inches(0.07)
    grid(slide, l + pad, t + title_h + Inches(0.02), w - 2 * pad, h - title_h - Inches(0.09),
         labels, cols, CORE_BG, CORE_LT, 8.8)


def module_row(slide, l, t, w, h, title, labels, title_w=Inches(1.9), chevron=False):
    card(slide, l, t, w, h, WHITE, CORE, 1.0, 0.08)
    text(slide, l + Inches(0.06), t, title_w - Inches(0.06), h, title,
         10, True, CORE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pad = Inches(0.07)
    x0 = l + title_w + Inches(0.04)
    iw = w - title_w - Inches(0.04) - pad
    if not chevron:
        grid(slide, x0, t + pad, iw, h - 2 * pad, labels, len(labels), CORE_BG, CORE_LT, 8.8)
        return
    n = len(labels)
    cw = iw / n
    for i, lab in enumerate(labels):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = shape(slide, kind, x0 + i * cw, t + pad, cw + Inches(0.05), h - 2 * pad,
                   CORE_BG, CORE_LT, 0.75, adj=0.22)
        label_shape(sh, lab, 8.8, False, INK)


def slide_main(prs):
    s = blank(prs)
    title_block(
        s, "设计态 AI 集群仿真器  ·  立项核心",
        "以「负载引擎」为核心的设计态仿真器：负载刻画准，系统设计才可信",
        "底层 NPU / 内存存储 / 网络仿真已具备；本项目补齐负载的画像、生成与导入，形成「真实负载 → 可外推负载 → 设计结论」闭环。")

    x0 = Inches(1.42)
    x1 = Inches(12.97)
    full = x1 - x0
    mid = x0 + full / 2

    # Layer 1: inputs
    t, h = Inches(1.2), Inches(0.46)
    layer_tag(s, t, h, "输入层", None, BLUE)
    grid(s, x0, t, full, h, [
        "生产 Trace（msprof / HCCL / 服务日志）",
        "模型结构与并行配置",
        "业务流量预测与 SLO 分级",
        "未来模型与集群规模假设",
    ], 4, BLUE_BG, None, 9.5)
    down_arrow(s, mid, t + h + Inches(0.03))

    # Layer 2: workload engine
    et, eh = Inches(1.86), Inches(3.78)
    layer_tag(s, et, eh, "负载引擎", "本项目核心", CORE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0, et, full, eh, CORE_BG, CORE, 1.75, 0.02)
    pad = Inches(0.08)
    cal_w = Inches(1.78)
    ix0 = x0 + pad
    iw = full - 2 * pad - cal_w - pad
    gap = Inches(0.08)
    mw = (iw - 2 * gap) / 3
    r1t, r1h = et + pad, Inches(1.82)
    module(s, ix0, r1t, mw, r1h, "① 基于计算图展开的模型负载画像", [
        "框架计算图抽取", "并行策略符号化展开",
        "多级算子代价模型", "通信量与访存量推导",
        "MoE 路由偏斜建模", "KV Cache 动态增长建模",
    ], 2)
    module(s, ix0 + mw + gap, r1t, mw, r1h, "② 基于随机过程的到达与请求建模", [
        "突发到达建模\nMMPP / Hawkes", "日周期与多租户叠加",
        "输入/输出长度\nCopula 联合建模", "会话与前缀复用建模",
        "训练作业规模与故障建模", "SLO 分级条件化生成",
    ], 2)
    module(s, ix0 + 2 * (mw + gap), r1t, mw, r1h, "③ 基于统计拟合的 Trace 重构与外推", [
        "多源 Trace 时钟对齐\n与依赖重建", "迭代周期性压缩",
        "混合分布拟合\n与 KS/W1 检验", "rank 对称规模外推\n1K → 100K",
        "跨硬件负载重定向", "未来模型 what-if 改写",
    ], 2)

    r2t, r2h = r1t + r1h + gap, Inches(0.84)
    module_row(s, ix0, r2t, iw, r2h, "④ 面向 DSE 的\n流水线化负载数据工厂", [
        "多源采集\n与脱敏", "负载特征化\n抽取", "参数化 / 生成式\n负载合成",
        "分布 + 效用\n保真校验", "场景化负载库\n版本管理",
    ], chevron=True)

    r3t = r2t + r2h + gap
    r3h = et + eh - pad - r3t
    module_row(s, ix0, r3t, iw, r3h, "⑤ 统一负载 IR\n与多保真导入", [
        "硬件无关执行图 IR\n兼容 Chakra ET", "L0 解析 / L1 算子图\n/ L2 包级切换",
        "集合通信\n分解为网络流", "访存与 KV 流量\n映射内存仿真", "代表迭代\n采样加速",
    ])

    cx = ix0 + iw + pad
    ct, ch = et + pad, eh - 2 * pad
    card(s, cx, ct, cal_w, ch, WHITE, CORE, 1.0, 0.04)
    text(s, cx + Inches(0.04), ct + Inches(0.03), cal_w - Inches(0.08), Inches(0.5),
         "⑥ 基于实测对标的\n闭环校准", 10, True, CORE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    grid(s, cx + Inches(0.07), ct + Inches(0.58), cal_w - Inches(0.14), ch - Inches(0.65), [
        "小规模集群实测对标", "分层误差归因", "代价模型系数回灌",
        "分布保真评估", "效用保真评估",
    ], 1, CORE_BG, CORE_LT, 8.8)

    down_arrow(s, ix0 + iw / 2, et + eh + Inches(0.03))
    shape(s, MSO_SHAPE.UP_ARROW, cx + cal_w / 2 - Inches(0.16), et + eh + Inches(0.03),
          Inches(0.32), Inches(0.14), CORE, adj=0.5)

    # Layer 3: simulation kernels
    kt, kh = et + eh + Inches(0.2), Inches(0.46)
    layer_tag(s, kt, kh, "仿真内核", "已具备", MUTED)
    grid(s, x0, kt, full, kh, ["NPU 仿真", "内存 / 存储仿真", "网络仿真"], 3, GRAY_BG, LINE, 10, 0.08)
    down_arrow(s, mid, kt + kh + Inches(0.03))

    # Layer 4: design outputs
    ot, oh = kt + kh + Inches(0.2), Inches(0.46)
    layer_tag(s, ot, oh, "设计输出", "DSE", GREEN)
    grid(s, x0, ot, full, oh, [
        "超节点规模与组网拓扑", "Scale-up / Scale-out 带宽配比", "HBM / 池化内存 / 存储层级容量",
        "PD 配比与并行策略", "瓶颈定位与 TCO / 能效评估",
    ], 5, GREEN_BG, None, 9.5, 0.08)

    footer(s, 1, "设计态 AI 集群仿真器  ·  负载引擎立项  ·  内部讨论稿")
    s.notes_slide.notes_text_frame.text = (
        "讲解主线：底层算力/内存/网络仿真解决的是「系统怎么跑」，而设计态真正的输入是「跑什么」。"
        "设计态的目标系统和未来模型往往还不存在，所以负载不能只回放历史 Trace，必须可画像、可参数化、可外推。\n"
        "①通过计算图展开刻画单作业内部的计算/通信/访存结构；②通过随机过程刻画多作业、多请求在时间上的叠加；"
        "③通过统计拟合把真实生产 Trace 变成可放大、可迁移的负载；④把以上能力流水线化、资产化，批量产出 DSE 输入；"
        "⑤统一 IR 让同一份负载以不同保真度驱动已有仿真内核；⑥仿真结果与实测对标，误差回灌，保证可信度。")
    return s


def slide_detail(prs):
    s = blank(prs)
    title_block(
        s, "备份页  ·  核心技术展开",
        "负载引擎六项核心技术：关键问题、技术路线、业界参考与难点",
        "按「单作业结构 → 多作业时间叠加 → 真实数据拟合 → 流水线资产化 → 导入仿真 → 闭环校准」组织。")

    data = [
        ["核心技术", "要回答的关键问题", "技术路线", "业界参考", "难点 / 风险"],
        ["① 模型画像",
         "一个作业在给定并行策略下，每个 rank 何时算什么、传多少、访存多少？",
         "计算图抽取（框架图 dump/ET）→ 并行策略符号化展开 → 算子代价模型（Roofline+实测表+回归）→ MoE 路由/KV 动态注入",
         "Chakra ET、Calculon、Vidur 算子 Profiling、Lumos",
         "未见 shape 与新算子的代价外推；MoE 路由偏斜的刻画"],
        ["② 到达与请求分布",
         "多请求/多作业在时间上如何叠加？突发、周期、长尾如何刻画？",
         "MMPP/Hawkes 刻画突发；输入/输出长度 Copula 联合建模；会话/前缀复用建模；训练作业规模幂律与故障模型",
         "Azure LLM Trace、BurstGPT、Mooncake Trace、ServeGen",
         "生产数据可得性；Agent/RL 等新形态负载缺少历史样本"],
        ["③ Trace 拟合与外推",
         "如何把小规模、旧硬件上的真实 Trace 变成大规模、新硬件的负载？",
         "多源对齐与依赖重建 → 迭代去重压缩 → 分布拟合 + 检验 → rank 对称扩展、跨硬件重定向、模型 what-if 改写",
         "Chakra trace link/converter、SimAI、Echo",
         "外推非线性（拥塞、尾延迟）；跨硬件重定向时依赖关系变化"],
        ["④ 负载数据工厂",
         "如何让负载可复用、可追溯、可批量产出？",
         "采集→清洗脱敏→特征化→建模→合成→校验→入库流水线；参数化生成器 + 生成式合成；负载版本与血缘管理",
         "Mystique（Meta）、MLPerf 负载集",
         "脱敏后的保真度；生成式模型的可解释性"],
        ["⑤ 统一负载 IR 与导入",
         "同一份负载如何以不同保真度驱动 NPU/内存/网络内核？",
         "硬件无关执行图 IR（兼容 Chakra ET）；L0/L1/L2 多保真切换；集合通信分解为流；代表迭代采样加速",
         "ASTRA-sim 2.0、SimAI（ns-3 后端）",
         "多内核协同的时间同步；万卡级图规模与仿真速度"],
        ["⑥ 闭环校准",
         "仿真结论有多可信？误差来自哪里？",
         "小规模实测对标 → 分层误差归因（算子/通信/调度）→ 系数回灌；分布保真 + 效用保真双指标",
         "ASTRA-sim/SimAI 实测验证方法",
         "大规模实测资源有限；误差随规模放大"],
    ]
    l, t = Inches(0.36), Inches(1.25)
    col_w = [Inches(1.7), Inches(2.5), Inches(4.05), Inches(2.2), Inches(2.16)]
    ts = s.shapes.add_table(len(data), len(data[0]), l, t, sum(col_w, Inches(0)), Inches(4.74))
    table = ts.table
    for i, cw in enumerate(col_w):
        table.columns[i].width = cw
    table.rows[0].height = Inches(0.36)
    for r in range(1, len(data)):
        table.rows[r].height = Inches(0.73)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = HEAD
            elif c == 0:
                cell.fill.fore_color.rgb = CORE_LT
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else ROW_ALT
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = val
            if r == 0:
                _style(run, 10.5, True, WHITE)
            elif c == 0:
                _style(run, 10.5, True, CORE)
            else:
                _style(run, 9.2, False, INK)

    sy, sh_ = Inches(6.22), Inches(0.66)
    rect(s, Inches(0.36), sy, Inches(12.61), sh_, WHITE, LINE, 0.75)
    rect(s, Inches(0.36), sy, Inches(1.3), sh_, BLUE)
    text(s, Inches(0.36), sy, Inches(1.3), sh_,
         [[("立项目标", 11, True, WHITE)], [("建议值，待对齐", 8, False, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    kpis = [
        ("≤10%", "迭代时间/吞吐仿真误差（千卡实测对标）"),
        ("1K→100K", "千卡 Trace 外推生成十万卡负载"),
        ("6类·30+", "场景类别 · 版本化标准负载"),
        ("≥1000", "配置/轮批量 DSE；解析级分钟、包级小时"),
    ]
    kx0 = Inches(1.74)
    kwid = (Inches(12.97) - kx0) / len(kpis)
    for i, (num, desc) in enumerate(kpis):
        x = kx0 + i * kwid
        if i:
            rect(s, x - Inches(0.02), sy + Inches(0.12), Inches(0.012), sh_ - Inches(0.24), LINE)
        text(s, x, sy, Inches(1.32), sh_, num, 13, True, CORE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(1.3), sy, kwid - Inches(1.34), sh_, desc, 8.8, False, INK,
             anchor=MSO_ANCHOR.MIDDLE)

    footer(s, 2, "参考：ASTRA-sim 2.0/Chakra、SimAI（NSDI'25）、Vidur（MLSys'24）、NVIDIA Omniverse DSX、"
                 "Mystique（ISCA'23）、ServeGen、BurstGPT、Mooncake")
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_main(prs)
    slide_detail(prs)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
