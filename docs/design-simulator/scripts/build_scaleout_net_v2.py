#!/usr/bin/env python3
"""ScaleOut network simulation charter slide v2.

v1 on main is the shallow noun list and stays archived.
v2 is the confirmed, implementable cut: discrete-event scheduling on
CLOS/Dragonfly, link-failure injection, control-plane reconvergence and
forwarding-path check; RoCE vs UBoE on send/completion queues, ECN/backpressure
and timeout/selective retransmit, with collectives replayed in event order;
182X DMA descriptors, doorbells and queues, hardware offload, then endpoint
latency/throughput calibration.

Same chrome as framework DAG v6. No gray「已具备」layer.
Boxes hold technique names only; order sits in captions and the right column.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / ".cursor/skills/pptx-layered-arch-diagram"))

from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches  # noqa: E402

from arch_kit import (  # noqa: E402
    blank,
    footer as kit_footer,
    label,
    new_deck,
    rbox,
    rect,
    shape,
    text,
)

OUT = ROOT / "docs/design-simulator/设计态仿真器_ScaleOut网络仿真能力_v2.pptx"

# AICO-PPT: brand red marks the page focus; slate marks structure.
RED = RGBColor(0xB5, 0x33, 0x3B)
RED_SOFT = RGBColor(0xFD, 0xF0, 0xF1)
RED_LINE = RGBColor(0xE0, 0xA3, 0xA7)
SLATE = RGBColor(0x56, 0x64, 0x72)
ARROW_GREY = RGBColor(0xC8, 0xC8, 0xCC)
INK = RGBColor(0x1A, 0x1A, 0x1C)
BODY = RGBColor(0x58, 0x58, 0x60)
LINE = RGBColor(0xE7, 0xE7, 0xE7)
PAPER = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DIAG_LEFT, DIAG_RIGHT = 0.22, 6.56
DIAG_TOP, DIAG_BOTTOM = 2.16, 7.14
TAG_W = 0.68
ARROW = 0.11
CAP_H = 0.20
CAP_GAP = 0.06
INPUT_H = 0.64
OUTPUT_H = 0.72

# Names are mechanisms a simulator can actually build. Order is outside the boxes.
INPUTS = ["拓扑与链路故障", "集合通信事件", "协议与拥塞参数", "182X 队列参数"]
TOPO = ["离散事件调度", "CLOS/Dragonfly", "链路故障注入", "控制面重收敛", "转发路径核对"]
PROTO = ["发送与完成队列", "ECN/拥塞反压", "超时与选择重传", "集合通信\n事件序回放"]
NIC = ["182X DMA\n描述符", "门铃与队列", "硬件 offload", "端侧延迟\n吞吐校准"]
OUTPUTS = ["重构路径时序", "集合完成时间", "RoCE/UBoE\n时序差", "端侧延迟吞吐"]

# cards: names side by side. Topology has five names because each one is a separate mechanism.
LAYERS = [
    ("① 事件驱动的拓扑仿真", "在 CLOS/Dragonfly 上调度事件，故障后核对转发路径", "cards", TOPO),
    ("② RoCE / UBoE 行为对比", "同一流量对比 RoCE / UBoE，集合通信按事件序回放", "cards", PROTO),
    ("③ 182X 端侧行为校准", "DMA 描述符、门铃与队列、硬件 offload，再校准端侧", "cards", NIC),
]

BANNER = [
    ("面向场景",
     "面向 10T～50T+、约 50 万卡 ScaleOut，在设计态做事件级网络仿真，对照吞吐、延迟、拥塞与瓶颈"),
    ("竞争力目标",
     "同一流量对比 RoCE 与 UBoE 的队列、ECN/拥塞反压与重传，并以 182X 的 DMA 描述符、门铃与硬件 offload 校准端侧延迟与吞吐"),
    ("技术挑战",
     "离散事件在五十万卡的 CLOS/Dragonfly 上可推进，故障后控制面可重收敛，集合通信按事件序可复现"),
]

TECHS = [
    ("关键技术 1", "离散事件调度与转发路径核对", [
        ("事件推进：", "在 CLOS/Dragonfly 上做离散事件调度，再注入链路故障"),
        ("路径核对：", "控制面重收敛后，逐跳核对转发路径"),
    ]),
    ("关键技术 2", "RoCE 与 UBoE 行为对比", [
        ("行为对比：", (
            "同一流量分别比较发送与完成队列、",
            "ECN/拥塞反压，以及超时与选择重传",
        )),
        ("事件回放：", "集合通信按事件序回放，对照完成时序"),
    ]),
    ("关键技术 3", "182X DMA、门铃与硬件 offload", [
        ("端侧建模：", (
            "用 182X DMA 描述符、门铃与队列、",
            "硬件 offload 描述端侧收发",
        )),
        ("延迟校准：", "用实测数据校准端侧延迟与吞吐"),
    ]),
]


def _down_arrow(slide, cx, y, h=0.10) -> None:
    shape(
        slide, MSO_SHAPE.DOWN_ARROW,
        Inches(cx - 0.10), Inches(y), Inches(0.20), Inches(h),
        ARROW_GREY, adj=0.5,
    )


def _tag(slide, x, y, h, name, sub, fill) -> None:
    rect(slide, Inches(x), Inches(y), Inches(TAG_W), Inches(h), fill)
    lines = [[(part, 10, True, WHITE)] for part in name.split("\n")]
    if sub:
        lines.append([(sub, 8, False, WHITE)])
    text(slide, Inches(x), Inches(y), Inches(TAG_W), Inches(h), lines,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _tighten(sh) -> None:
    for p in sh.text_frame.paragraphs:
        p.space_after = 0


def _cards(slide, x, y, w, h, items, size=10) -> None:
    gap = 0.06
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, lab in enumerate(items):
        sh = rbox(
            slide, Inches(x + i * (cw + gap)), Inches(y), Inches(cw), Inches(h),
            WHITE, LINE, 0.9, 0.12,
        )
        label(sh, lab, size, False, INK)
        _tighten(sh)


def _flow(slide, x, y, w, h, items, size=10) -> None:
    """Ordered steps. Later chevrons cover the previous tip."""
    n = len(items)
    overlap = 0.04
    cw = (w - overlap) / n
    for i, lab in enumerate(items):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        extra = overlap if i < n - 1 else 0.0
        sh = shape(
            slide, kind,
            Inches(x + i * cw), Inches(y), Inches(cw + extra), Inches(h),
            WHITE, LINE, 0.9, adj=0.14,
        )
        label(sh, lab, size, False, INK)
        _tighten(sh)


def _caption(slide, x, y, w, h, body) -> None:
    text(slide, Inches(x), Inches(y), Inches(w), Inches(h), body,
         10, False, BODY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def _title(slide, x, y, w, h, body) -> None:
    text(slide, Inches(x), Inches(y), Inches(w), Inches(h), body,
         12, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def _module(slide, x, y, w, h, title, caption, kind, items) -> None:
    rbox(slide, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, LINE, 0.9, 0.05)
    head_h, cap_h = 0.24, 0.18
    _title(slide, x + 0.08, y + 0.02, w - 0.16, head_h, title)
    _caption(slide, x + 0.08, y + head_h, w - 0.16, cap_h, caption)
    body_y = y + head_h + cap_h + 0.02
    body_h = h - (body_y - y) - 0.06
    if body_h < 0.34:
        raise RuntimeError(f"module body too short: {body_h:.3f} in")
    draw = _flow if kind == "flow" else _cards
    # Five names share one row, so those labels run slightly smaller.
    size = 9 if len(items) >= 5 else 10
    draw(slide, x + 0.08, body_y, w - 0.16, body_h, items, size)


def _core(slide, x, y, w, h) -> None:
    rbox(slide, Inches(x), Inches(y), Inches(w), Inches(h), RED_SOFT, RED_LINE, 1.0, 0.03)
    pad_x, pad_y = 0.08, 0.05
    ix, iw = x + pad_x, w - 2 * pad_x
    cap_h = 0.20
    arrow_h = 0.14
    _caption(slide, ix, y + pad_y, iw, cap_h, "自上而下三层：事件推进、协议行为、端侧校准")
    inner_top = y + pad_y + cap_h + CAP_GAP
    inner_bot = y + h - pad_y
    module_h = (inner_bot - inner_top - 2 * arrow_h) / 3
    if module_h < 0.78:
        raise RuntimeError(f"ScaleOut core module too short: {module_h:.3f} in")
    cursor = inner_top
    mid = x + w / 2
    for i, (title, caption, kind, items) in enumerate(LAYERS):
        _module(slide, ix, cursor, iw, module_h, title, caption, kind, items)
        cursor += module_h
        if i < len(LAYERS) - 1:
            _down_arrow(slide, mid, cursor + 0.02, arrow_h - 0.04)
            cursor += arrow_h


def _io(slide, x, y, w, h, caption, items) -> None:
    _caption(slide, x, y, w, CAP_H, caption)
    _cards(slide, x, y + CAP_H + CAP_GAP, w, h - CAP_H - CAP_GAP, items, 10)


def _diagram(slide) -> None:
    tag_x = DIAG_LEFT
    x = tag_x + TAG_W + 0.05
    w = DIAG_RIGHT - x
    core_h = (DIAG_BOTTOM - DIAG_TOP) - INPUT_H - OUTPUT_H - 2 * ARROW
    y = DIAG_TOP
    _tag(slide, tag_x, y, INPUT_H, "输入层", None, SLATE)
    _io(slide, x, y, w, INPUT_H, "四类输入同时抽出", INPUTS)
    y += INPUT_H
    _down_arrow(slide, x + w / 2, y + 0.005, ARROW - 0.01)
    y += ARROW
    _tag(slide, tag_x, y, core_h, "仿真\n技术", "本页核心", RED)
    _core(slide, x, y, w, core_h)
    y += core_h
    _down_arrow(slide, x + w / 2, y + 0.005, ARROW - 0.01)
    y += ARROW
    _tag(slide, tag_x, y, OUTPUT_H, "输出层", None, SLATE)
    _io(slide, x, y, w, OUTPUT_H, "从事件结果分出四类对照", OUTPUTS)


def _banner(slide) -> None:
    top, height = 0.78, 1.08
    rbox(slide, Inches(0.22), Inches(top), Inches(12.88), Inches(height), WHITE, LINE, 0.9, 0.05)
    row_h = height / len(BANNER)
    for i, (name, body) in enumerate(BANNER):
        yy = top + i * row_h
        rect(slide, Inches(0.36), Inches(yy + 0.10), Inches(0.07), Inches(row_h - 0.20), RED)
        text(slide, Inches(0.52), Inches(yy), Inches(1.22), Inches(row_h), name,
             12, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        text(slide, Inches(1.78), Inches(yy), Inches(11.10), Inches(row_h), body,
             12, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def _tech_column(slide) -> None:
    left, right = 6.72, 13.08
    top, bottom = DIAG_TOP, DIAG_BOTTOM
    width = right - left
    gap = 0.07
    card_h = (bottom - top - gap * (len(TECHS) - 1)) / len(TECHS)
    text(slide, Inches(left), Inches(top - 0.22), Inches(width), Inches(0.20),
         "关键技术", 12, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    for i, (kicker, title, points) in enumerate(TECHS):
        yy = top + i * (card_h + gap)
        rbox(slide, Inches(left), Inches(yy), Inches(width), Inches(card_h), WHITE, LINE, 0.9, 0.06)
        rect(slide, Inches(left), Inches(yy + 0.08), Inches(0.07), Inches(card_h - 0.16), RED)
        text(slide, Inches(left + 0.18), Inches(yy + 0.05), Inches(width - 0.30), Inches(0.22),
             kicker, 11, True, RED)
        text(slide, Inches(left + 0.18), Inches(yy + 0.26), Inches(width - 0.30), Inches(0.28),
             title, 13, True, INK)
        slot_top = yy + 0.56
        slot_h = (card_h - 0.64) / 2
        for n, (lead, desc) in enumerate(points):
            parts = desc if isinstance(desc, tuple) else (desc,)
            lines = [[(f"{n + 1}）{lead}", 11, True, INK), (parts[0], 11, False, BODY)]]
            lines += [[(part, 11, False, BODY)] for part in parts[1:]]
            text(
                slide,
                Inches(left + 0.18),
                Inches(slot_top + n * slot_h),
                Inches(width - 0.32),
                Inches(slot_h),
                lines,
                anchor=MSO_ANCHOR.TOP,
            )


def build() -> None:
    prs = new_deck()
    slide = blank(prs)
    rect(slide, 0, 0, prs.slide_width, prs.slide_height, PAPER)
    rect(slide, 0, 0, prs.slide_width, Inches(0.06), RED)
    text(slide, Inches(0.28), Inches(0.12), Inches(12.6), Inches(0.24),
         "设计态 AI 集群仿真器  ·  ScaleOut 网络仿真", 12, True, RED)
    text(slide, Inches(0.28), Inches(0.36), Inches(12.7), Inches(0.38),
         "以事件驱动仿真拓扑、协议与模组，支撑十万～五十万卡 ScaleOut 网络设计决策",
         20, True, INK)
    _banner(slide)
    text(slide, Inches(DIAG_LEFT), Inches(DIAG_TOP - 0.22), Inches(3.4), Inches(0.20),
         "能力结构", 12, True, SLATE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    _diagram(slide)
    _tech_column(slide)
    kit_footer(slide, "ScaleOut 网络仿真  ·  立项讨论稿", 1, 1)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), RED)
    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "v2 为已确认的可落地机制版。v1 浅稿归档。版式对齐框架侧 DAG v6，不画灰色「已具备」层。\n"
        "面向 10T～50T+、约 50 万卡 ScaleOut。标题中的十万～五十万卡是离散事件要推进的规模。\n"
        "输入四类同时抽出：拓扑与链路故障、集合通信事件、协议与拥塞参数、182X 队列参数。\n"
        "① 事件驱动的拓扑仿真。在 CLOS/Dragonfly 上做离散事件调度，注入链路故障，"
        "跑控制面重收敛，再逐跳核对转发路径。这些都是现成的网络仿真步骤，不另起一套名字。\n"
        "② RoCE 与 UBoE 行为对比。同一流量分别比较发送与完成队列、ECN/拥塞反压、超时与选择重传。"
        "集合通信按事件序回放，对照完成时序，不按平均带宽估算。\n"
        "③ 182X 端侧行为校准。用 DMA 描述符、门铃与队列、硬件 offload 描述收发，"
        "再用实测数据校准端侧延迟与吞吐。换模组不改上层拓扑与协议的事件接口。\n"
        "输出：重构路径时序、集合完成时间、RoCE/UBoE 时序差、端侧延迟吞吐。\n"
        "配色：品牌红 #b5333b 只标仿真技术层；输入输出用灰蓝 #566472；过渡箭头浅灰 #c8c8cc。"
        "框内只写技术点名称，先后写在框外短注和右栏。"
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
