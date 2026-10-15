#!/usr/bin/env python3
"""ScaleOut network simulation charter slide v3.

Layout matches框架侧 DAG v6 / ScaleOut v2: brow + title, 3 banner rows,
left layered arch (problems strip → input → core → output, NO gray 已具备),
right 3 key-tech cards.

v3 vs v2 (WeChat feedback):
- Key-tech titles = deep nouns (基于…), not verb-object phrases.
- Core module titles = 方法+对象 (基于…的…仿真).
- ~4 noun boxes per core layer; no component piles (队列/功能堆砌).
- Still no 「校准」; fidelity/modeling phrasing only.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / ".cursor/skills/pptx-layered-arch-diagram"))

from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches  # noqa: E402

from arch_kit import (  # noqa: E402
    blank,
    footer as kit_footer,
    label,
    new_deck,
    rbox,
    rect,
    shape,
    text,
)

OUT = ROOT / "docs/design-simulator/设计态仿真器_ScaleOut网络仿真能力_v3.pptx"

# AICO-PPT: brand red marks the page focus; slate marks structure.
RED = RGBColor(0xB5, 0x33, 0x3B)
RED_SOFT = RGBColor(0xFD, 0xF0, 0xF1)
RED_LINE = RGBColor(0xE0, 0xA3, 0xA7)
SLATE = RGBColor(0x56, 0x64, 0x72)
ARROW_GREY = RGBColor(0xC8, 0xC8, 0xCC)
INK = RGBColor(0x1A, 0x1A, 0x1C)
BODY = RGBColor(0x58, 0x58, 0x60)
LINE = RGBColor(0xE7, 0xE7, 0xE7)
PAPER = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHIP_FILL = RGBColor(0xF5, 0xF5, 0xF6)

DIAG_LEFT, DIAG_RIGHT = 0.22, 6.56
# Same geometry as v2 (problems strip between column titles and stack).
DIAG_TOP, DIAG_BOTTOM = 2.48, 7.14
TAG_W = 0.68
ARROW = 0.09
INPUT_H, OUTPUT_H = 0.32, 0.34

PROBLEMS = [
    "ScaleOut\n带宽设计",
    "收敛比分析",
    "拥塞/丢包\n性能测算",
    "协议路径\n吞吐对照",
    "故障重构影响",
]

INPUTS = ["集群拓扑意图", "并行与通信模式", "协议与拥塞策略", "网卡与模组参数"]
OUTPUTS = ["带宽延迟对照", "集合通信效率", "瓶颈定位图", "RoCE/UBoE/\n网卡方案对比"]

CORE_MODULES = [
    (
        "① 基于拓扑意图的路由仿真",
        [
            "CLOS/Dragonfly\n拓扑建模",
            "ECMP\n多路径选路",
            "故障重构\n路径推演",
            "跨域收敛比\n分析",
        ],
    ),
    (
        "② 基于双协议的流量仿真",
        [
            "RoCE\n行为模型",
            "UBoE\n行为模型",
            "集合通信\n事件重放",
            "PFC/ECN\n拥塞策略",
        ],
    ),
    (
        "③ 基于 182X 的设备仿真",
        [
            "182X DMA\n行为模型",
            "硬件\noffload 模型",
            "网侧交换机\n模型",
            "模组可替换\n接口",
        ],
    ),
]

NOTE_INPUT = "四类输入同时抽出"
NOTE_CORE = "自上而下三层，协议层必须能切 RoCE / UBoE"
NOTE_OUTPUT = "从执行结果分出四类对照"
NOTE_PROBLEMS = "待验证 / 待测算"

BANNER = [
    (
        "面向场景",
        "面向 10T～50T+ 大模型、约 50 万卡量级集群，在设计态仿真 ScaleOut 网络方案与网卡行为，对照吞吐、延迟与瓶颈",
    ),
    (
        "竞争力目标",
        "可替换对比 RoCE 与 UBoE 等协议路径，并纳入 182X 网卡级行为，给出可量化的方案取舍依据",
    ),
    (
        "技术挑战",
        "超大规模拓扑与路由保真、集合通信与流量跨 ScaleOut 可复现、端侧网卡与网侧交换机行为可插拔",
    ),
]

TECHS = [
    (
        "关键技术 1",
        "基于可扩拓扑的多路径交叉验证",
        [
            ("规模化：", "先建可扩拓扑，再注入路由与故障路径"),
            ("可对照：", "同一负载下对比不同拓扑/路由配置的瓶颈位置"),
        ],
    ),
    (
        "关键技术 2",
        "基于双协议的路径对照仿真",
        [
            ("可切换：", "同一流量意图分别走 RoCE / UBoE 行为模型"),
            ("可复现：", "集合通信与跨 ScaleOut 流量按事件序回放"),
        ],
    ),
    (
        "关键技术 3",
        "基于 182X 的模组保真仿真",
        [
            (
                "保真：",
                (
                    "用 182X DMA/offload 与交换机行为",
                    "建模端侧与网侧延迟、吞吐",
                ),
            ),
            ("可替换：", "设备接口固定，换网卡/交换机模型不改上层拓扑与协议层"),
        ],
    ),
]


def _down_arrow(slide, cx, y, h=0.10) -> None:
    shape(
        slide,
        MSO_SHAPE.DOWN_ARROW,
        Inches(cx - 0.10),
        Inches(y),
        Inches(0.20),
        Inches(h),
        ARROW_GREY,
        adj=0.5,
    )


def _tag(slide, x, y, h, name, sub, fill) -> None:
    rect(slide, Inches(x), Inches(y), Inches(TAG_W), Inches(h), fill)
    lines = [[(part, 10, True, WHITE)] for part in name.split("\n")]
    if sub:
        lines.append([(sub, 8, False, WHITE)])
    text(
        slide,
        Inches(x),
        Inches(y),
        Inches(TAG_W),
        Inches(h),
        lines,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def _cards(slide, x, y, w, h, items, size=9) -> None:
    gap = 0.04
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, lab in enumerate(items):
        sh = rbox(
            slide,
            Inches(x + i * (cw + gap)),
            Inches(y),
            Inches(cw),
            Inches(h),
            WHITE,
            LINE,
            0.9,
            0.10,
        )
        if "\n" in lab:
            label(
                sh,
                [[(part, size - 0.5, False, INK)] for part in lab.split("\n")],
            )
            for p in sh.text_frame.paragraphs:
                p.space_after = 0
        else:
            label(sh, lab, size, False, INK)


def _problems_strip(slide) -> None:
    """待验证问题 chips between column titles and the input→core→output stack."""
    tag_x = DIAG_LEFT
    x = tag_x + TAG_W + 0.05
    w = DIAG_RIGHT - x
    strip_h = 0.36
    y = DIAG_TOP - strip_h - 0.06
    text(
        slide,
        Inches(tag_x),
        Inches(y),
        Inches(TAG_W),
        Inches(strip_h),
        [[("待验", 9, True, SLATE)], [("证项", 9, True, SLATE)]],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    gap = 0.04
    n = len(PROBLEMS)
    cw = (w - gap * (n - 1)) / n
    for i, lab in enumerate(PROBLEMS):
        sh = rbox(
            slide,
            Inches(x + i * (cw + gap)),
            Inches(y + 0.02),
            Inches(cw),
            Inches(strip_h - 0.04),
            CHIP_FILL,
            LINE,
            0.85,
            0.12,
        )
        if "\n" in lab:
            label(
                sh,
                [[(part, 8.5, False, INK)] for part in lab.split("\n")],
            )
            for p in sh.text_frame.paragraphs:
                p.space_after = 0
        else:
            label(sh, lab, 8.5, False, INK)


def _core(slide, x, y, w, h) -> None:
    rbox(slide, Inches(x), Inches(y), Inches(w), Inches(h), RED_SOFT, RED_LINE, 1.0, 0.03)
    pad_x, pad_y = 0.06, 0.04
    ix, iw = x + pad_x, w - 2 * pad_x
    note_h = 0.18
    text(
        slide,
        Inches(ix),
        Inches(y + pad_y),
        Inches(iw),
        Inches(note_h),
        NOTE_CORE,
        8.5,
        False,
        BODY,
        PP_ALIGN.LEFT,
        MSO_ANCHOR.MIDDLE,
    )
    n = len(CORE_MODULES)
    gap = 0.05
    body_top = y + pad_y + note_h + 0.01
    body_h = h - (body_top - y) - pad_y
    mod_h = (body_h - gap * (n - 1)) / n
    title_h = 0.18
    # 4 boxes/layer: slightly taller cards than v2's 6-box packing
    card_h = min(mod_h - title_h - 0.03, 0.52)
    for i, (title, items) in enumerate(CORE_MODULES):
        yy = body_top + i * (mod_h + gap)
        text(
            slide,
            Inches(ix),
            Inches(yy),
            Inches(iw),
            Inches(title_h),
            title,
            10.5,
            True,
            RED,
            PP_ALIGN.LEFT,
            MSO_ANCHOR.MIDDLE,
        )
        _cards(slide, ix, yy + title_h + 0.01, iw, card_h, items, 8.5)
        if i < n - 1:
            _down_arrow(
                slide,
                x + w / 2,
                yy + mod_h + 0.002,
                max(gap - 0.008, 0.04),
            )


def _diagram(slide) -> None:
    tag_x = DIAG_LEFT
    x = tag_x + TAG_W + 0.05
    w = DIAG_RIGHT - x
    core_h = (DIAG_BOTTOM - DIAG_TOP) - INPUT_H - OUTPUT_H - 2 * ARROW
    y = DIAG_TOP

    _tag(slide, tag_x, y, INPUT_H, "输入层", None, SLATE)
    _cards(slide, x, y, w, INPUT_H, INPUTS, 9.5)
    y += INPUT_H
    _down_arrow(slide, x + w / 2, y + 0.004, ARROW - 0.01)
    y += ARROW

    _tag(slide, tag_x, y, core_h, "仿真\n技术", "本页核心", RED)
    _core(slide, x, y, w, core_h)
    y += core_h
    _down_arrow(slide, x + w / 2, y + 0.004, ARROW - 0.01)
    y += ARROW

    text(
        slide,
        Inches(x),
        Inches(y - ARROW + 0.01),
        Inches(w * 0.72),
        Inches(ARROW - 0.02),
        NOTE_OUTPUT,
        8,
        False,
        BODY,
        PP_ALIGN.LEFT,
        MSO_ANCHOR.MIDDLE,
    )
    _tag(slide, tag_x, y, OUTPUT_H, "输出层", None, SLATE)
    _cards(slide, x, y, w, OUTPUT_H, OUTPUTS, 9)


def _banner(slide) -> None:
    top, height = 0.78, 1.08
    rbox(slide, Inches(0.22), Inches(top), Inches(12.88), Inches(height), WHITE, LINE, 0.9, 0.05)
    row_h = height / len(BANNER)
    for i, (name, body) in enumerate(BANNER):
        yy = top + i * row_h
        rect(slide, Inches(0.36), Inches(yy + 0.10), Inches(0.07), Inches(row_h - 0.20), RED)
        text(
            slide,
            Inches(0.52),
            Inches(yy),
            Inches(1.22),
            Inches(row_h),
            name,
            12,
            True,
            RED,
            PP_ALIGN.LEFT,
            MSO_ANCHOR.MIDDLE,
        )
        text(
            slide,
            Inches(1.78),
            Inches(yy),
            Inches(11.10),
            Inches(row_h),
            body,
            12,
            False,
            INK,
            PP_ALIGN.LEFT,
            MSO_ANCHOR.MIDDLE,
        )


def _tech_column(slide) -> None:
    left, right = 6.72, 13.08
    top, bottom = DIAG_TOP, DIAG_BOTTOM
    width = right - left
    gap = 0.07
    card_h = (bottom - top - gap * (len(TECHS) - 1)) / len(TECHS)
    text(
        slide,
        Inches(left),
        Inches(1.88),
        Inches(width),
        Inches(0.16),
        "关键技术",
        12,
        True,
        RED,
        PP_ALIGN.LEFT,
        MSO_ANCHOR.MIDDLE,
    )
    for i, (kicker, title, points) in enumerate(TECHS):
        yy = top + i * (card_h + gap)
        rbox(slide, Inches(left), Inches(yy), Inches(width), Inches(card_h), WHITE, LINE, 0.9, 0.06)
        rect(slide, Inches(left), Inches(yy + 0.08), Inches(0.07), Inches(card_h - 0.16), RED)
        text(
            slide,
            Inches(left + 0.18),
            Inches(yy + 0.05),
            Inches(width - 0.30),
            Inches(0.22),
            kicker,
            11,
            True,
            RED,
        )
        text(
            slide,
            Inches(left + 0.18),
            Inches(yy + 0.26),
            Inches(width - 0.30),
            Inches(0.26),
            title,
            13,
            True,
            INK,
        )
        slot_top = yy + 0.54
        slot_h = (card_h - 0.62) / 2
        for n, (lead, desc) in enumerate(points):
            parts = desc if isinstance(desc, tuple) else (desc,)
            lines = [[(f"{n + 1}）{lead}", 11, True, INK), (parts[0], 11, False, BODY)]]
            lines += [[(part, 11, False, BODY)] for part in parts[1:]]
            text(
                slide,
                Inches(left + 0.18),
                Inches(slot_top + n * slot_h),
                Inches(width - 0.32),
                Inches(slot_h),
                lines,
                anchor=MSO_ANCHOR.TOP,
            )


def build() -> None:
    prs = new_deck()
    slide = blank(prs)
    rect(slide, 0, 0, prs.slide_width, prs.slide_height, PAPER)
    rect(slide, 0, 0, prs.slide_width, Inches(0.06), RED)
    text(
        slide,
        Inches(0.28),
        Inches(0.12),
        Inches(12.6),
        Inches(0.24),
        "设计态 AI 集群仿真器  ·  ScaleOut 网络仿真",
        12,
        True,
        RED,
    )
    text(
        slide,
        Inches(0.28),
        Inches(0.36),
        Inches(12.7),
        Inches(0.38),
        "以拓扑、协议、模组三层仿真支撑十万～五十万卡 ScaleOut 网络设计决策",
        20,
        True,
        INK,
    )
    _banner(slide)
    text(
        slide,
        Inches(DIAG_LEFT),
        Inches(1.88),
        Inches(1.6),
        Inches(0.16),
        "能力结构",
        12,
        True,
        SLATE,
        PP_ALIGN.LEFT,
        MSO_ANCHOR.MIDDLE,
    )
    text(
        slide,
        Inches(DIAG_LEFT + 1.65),
        Inches(1.88),
        Inches(DIAG_RIGHT - DIAG_LEFT - 1.65),
        Inches(0.16),
        NOTE_INPUT,
        8.5,
        False,
        BODY,
        PP_ALIGN.RIGHT,
        MSO_ANCHOR.MIDDLE,
    )
    _problems_strip(slide)
    _diagram(slide)
    _tech_column(slide)
    kit_footer(slide, "ScaleOut 网络仿真能力  ·  立项讨论稿", 1, 1)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), RED)
    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "v3：相对 v2——关键技术改为有深度的名词（基于…）；核心模块标题改为方法+对象；"
        "每层约 4 个技术名词框，去掉组件/功能堆砌；仍无「校准」措辞。\n"
        "关键技术 1：基于可扩拓扑的多路径交叉验证\n"
        "关键技术 2：基于双协议的路径对照仿真\n"
        "关键技术 3：基于 182X 的模组保真仿真\n"
        "① 基于拓扑意图的路由仿真：CLOS/Dragonfly 拓扑建模、ECMP 多路径选路、故障重构路径推演、跨域收敛比分析。\n"
        "② 基于双协议的流量仿真：RoCE/UBoE 行为模型、集合通信事件重放、PFC/ECN 拥塞策略。\n"
        "③ 基于 182X 的设备仿真：182X DMA 行为模型、硬件 offload 模型、网侧交换机模型、模组可替换接口。\n"
        "配色遵循 AICO-PPT：品牌红 #b5333b 只标仿真技术核心，结构用灰蓝 #566472。"
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
