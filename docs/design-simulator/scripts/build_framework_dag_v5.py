#!/usr/bin/env python3
"""Framework DAG slide v5: v4 content; light-grey connector arrows, borderless red 难点 text."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / ".cursor/skills/pptx-layered-arch-diagram"))

from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches  # noqa: E402

from arch_kit import (  # noqa: E402
    blank,
    footer as kit_footer,
    label,
    new_deck,
    rbox,
    rect,
    shape,
    text,
)

OUT = ROOT / "docs/design-simulator/设计态仿真器_框架DAG仿真能力_v5.pptx"

# AICO-PPT: brand red marks the page focus; slate marks structure.
RED = RGBColor(0xB5, 0x33, 0x3B)
RED_SOFT = RGBColor(0xFD, 0xF0, 0xF1)
RED_LINE = RGBColor(0xE0, 0xA3, 0xA7)
SLATE = RGBColor(0x56, 0x64, 0x72)
ARROW_GREY = RGBColor(0xC8, 0xC8, 0xCC)  # connector arrows: light grey
INK = RGBColor(0x1A, 0x1A, 0x1C)
BODY = RGBColor(0x58, 0x58, 0x60)
LINE = RGBColor(0xE7, 0xE7, 0xE7)
PAPER = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FIELD = RGBColor(0xF5, 0xF5, 0xF6)

DIAG_LEFT, DIAG_RIGHT = 0.22, 6.56
DIAG_TOP, DIAG_BOTTOM = 2.16, 7.14
TAG_W = 0.68
ARROW = 0.11
INPUT_H, OUTPUT_H = 0.32, 0.34

# Rounded rectangles are nouns. Chevrons are ordered steps.
# Every box is (term, technique tag). Lane steps are (term, tag, tag).
# No sentence captions inside the diagram; difficulties are short 「难点」 chips.
SOURCES = [
    ("模型结构解析", "Config · Module 树"),
    ("源码静态抽取", "源码 AST 分析"),
    ("框架 IR 抽取", "FX · Jaxpr · MindIR"),
    ("运行 Trace 回填", "Kineto · msprof"),
]
BIND = [
    ("算子实例化", "算子分解映射"),
    ("形状精度绑定", "符号形状推导"),
    ("控制流展开", "子图内联"),
    ("多框架归一", "统一算子 IR"),
]
STAGE1_HARD = "难点：动态形状 · 图断点"

SCALE = [
    ("设备网格映射", "Device Mesh", "拓扑亲和放置"),
    ("rank 级展开", "SPMD 复制", "同构折叠"),
    ("并行策略切分", "切分传播", "集合通信插入"),
    ("批次并发扩展", "微批 1F1B", "连续批处理"),
]
SCALE_HARD = "难点：万卡图膨胀 · 切分一致"
SEQ = [
    ("序列并行切分", "SP / CP", "Ring Attn"),
    ("Prefix 复用", "KV 分页", "前缀树匹配"),
    ("量化算子替换", "W8A8 / FP8", "Q/DQ 插入"),
    ("融合与重计算", "融合 pass", "激活重计算"),
]
SEQ_HARD = "难点：变长序列 · KV 显存"

DOMAINS = [
    ("集合通信图", "AllReduce · A2A"),
    ("点对点流量图", "PP Send/Recv"),
    ("内存容量图", "HBM 生命周期"),
    ("数据读取图", "样本 · ckpt IO"),
]
SCHEDULE = [
    ("依赖展开", "跨 rank 依赖"),
    ("算子调度", "多流 · 通算重叠"),
    ("执行序编排", "约束拓扑排序"),
    ("多场景重放", "离散事件重放"),
]
STAGE3_HARD = "难点：通算重叠 · 跨 rank 同步"

INPUTS = ["训练框架", "推理框架", "RL 框架", "模型与源码"]
OUTPUTS = ["场景执行图", "网络流量图", "内存访存图", "吞吐瓶颈对照"]

BANNER = [
    ("面向场景",
     "面向训练、推理、RL 等框架负载的端到端行为仿真，对照不同规模、资源、调度与并行限制下的吞吐、排队及瓶颈"),
    ("竞争力目标",
     "从模型结构与源码导出统一 DAG，按场景动态变换出大卡执行图及网络、内存读取图，形成可复用的框架级泛化仿真"),
    ("技术挑战",
     "训练、推理、RL 框架迭代快，结构、源码与运行图粒度不一；规模、并行、序列、通信和访存一变，静态 DAG 难以统一表征"),
]

TECHS = [
    ("关键技术 1", "多源图捕获，再归一为可仿真 DAG", [
        ("多源捕获：", (
            "结构配置、源码 AST、框架 IR（FX / Jaxpr / MindIR）",
            "与 Profiler Trace 四路同时进入。",
        )),
        ("归一收束：", (
            "先分解、实例化算子，再推导符号形状与精度，",
            "内联控制流子图，归一为统一算子 IR。",
        )),
    ]),
    ("关键技术 2", "按场景做图变换，并派生域图", [
        ("双链变换：", (
            "规模链做 rank 级展开、切分传播与集合通信插入；",
            "序列链做 SP/CP 切分、KV 复用、量化替换与重计算。",
        )),
        ("执行图分叉：", (
            "双链汇合为 rank 级执行图，再派生集合通信、",
            "点对点流量、HBM 容量与数据读取四张域图。",
        )),
    ]),
    ("关键技术 3", "按约束调度，再做 Trace 重放", [
        ("调度顺序：", (
            "先补齐跨 rank 依赖，再做多流调度与通算重叠，",
            "最后按资源约束编排执行序。",
        )),
        ("重放产出：", (
            "离散事件重放多场景，输出执行图、流量图、访存图，",
            "以及吞吐、排队与瓶颈对照。",
        )),
    ]),
]


def _down_arrow(slide, cx, y, h=0.10) -> None:
    shape(
        slide, MSO_SHAPE.DOWN_ARROW,
        Inches(cx - 0.10), Inches(y), Inches(0.20), Inches(h),
        ARROW_GREY, adj=0.5,
    )


def _tag(slide, x, y, h, name, sub, fill) -> None:
    rect(slide, Inches(x), Inches(y), Inches(TAG_W), Inches(h), fill)
    lines = [[(part, 10, True, WHITE)] for part in name.split("\n")]
    if sub:
        lines.append([(sub, 8, False, WHITE)])
    text(slide, Inches(x), Inches(y), Inches(TAG_W), Inches(h), lines,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


TAG_GREY = RGBColor(0x80, 0x80, 0x88)


def _two_line(sh, term, tag, size=9.5, tag_size=7.5) -> None:
    label(sh, [[(term, size, False, INK)], [(tag, tag_size, False, TAG_GREY)]])
    for p in sh.text_frame.paragraphs:
        p.space_after = 0


def _cards(slide, x, y, w, h, items, size=10) -> None:
    """Noun boxes. items: str or (term, technique tag)."""
    gap = 0.06
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, lab in enumerate(items):
        sh = rbox(
            slide, Inches(x + i * (cw + gap)), Inches(y), Inches(cw), Inches(h),
            WHITE, LINE, 0.9, 0.12,
        )
        if isinstance(lab, tuple):
            _two_line(sh, *lab)
        else:
            label(sh, lab, size, False, INK)


def _flow(slide, x, y, w, h, items) -> None:
    """Ordered steps (term, tag). Later chevrons cover the previous tip, so the row reads left to right."""
    n = len(items)
    overlap = 0.04
    cw = (w - overlap) / n
    for i, (term, tag) in enumerate(items):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        extra = overlap if i < n - 1 else 0.0
        sh = shape(
            slide, kind,
            Inches(x + i * cw), Inches(y), Inches(cw + extra), Inches(h),
            WHITE, LINE, 0.9, adj=0.20,
        )
        _two_line(sh, term, tag, 9.5, 7.5)


def _hard(slide, right, y, h, body, size=8) -> None:
    """「难点」 as plain red text (no border, no fill), right-aligned at `right`."""
    w = 0.13 + 0.112 * sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in body) * size / 8
    text(slide, Inches(right - w), Inches(y), Inches(w), Inches(h), body,
         size, False, RED, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def _title(slide, x, y, w, h, body, hard=None) -> None:
    text(slide, Inches(x), Inches(y), Inches(w), Inches(h), body,
         11, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    if hard:
        _hard(slide, x + w, y + 0.015, h - 0.03, hard)


def _lane(slide, x, y, w, h, kicker, hard, items) -> None:
    """One downward chain inside its own card: step box + two technique tags per row."""
    rbox(slide, Inches(x), Inches(y), Inches(w), Inches(h), WHITE, LINE, 0.9, 0.06)
    head_h = 0.22
    text(
        slide, Inches(x + 0.08), Inches(y + 0.03), Inches(1.0), Inches(head_h - 0.02),
        kicker, 10, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )
    _hard(slide, x + w - 0.08, y + 0.045, head_h - 0.05, hard, 7.5)
    pad = 0.08
    inner_x, inner_w = x + pad, w - 2 * pad
    step_w, tag_gap = 0.90, 0.035
    tag_w = (inner_w - step_w - 2 * tag_gap) / 2
    top = y + head_h + 0.06
    bottom = y + h - 0.06
    gaps = len(items) - 1
    arrow_h = 0.07
    box_h = (bottom - top - gaps * arrow_h) / len(items)
    yy = top
    for i, (term, *tags) in enumerate(items):
        sh = rbox(
            slide, Inches(inner_x), Inches(yy), Inches(step_w), Inches(box_h),
            FIELD, LINE, 0.8, 0.18,
        )
        label(sh, term, 9.5, False, INK)
        for k, tag in enumerate(tags):
            tx = inner_x + step_w + tag_gap + k * (tag_w + tag_gap)
            tsh = rbox(
                slide, Inches(tx), Inches(yy + 0.025), Inches(tag_w), Inches(box_h - 0.05),
                WHITE, LINE, 0.75, 0.30,
            )
            label(tsh, tag, 7.5, False, BODY)
        if i < gaps:
            _down_arrow(slide, inner_x + step_w / 2, yy + box_h + 0.006, arrow_h - 0.012)
        yy += box_h
        if i < gaps:
            yy += arrow_h


def _fork(slide, x, y, w, h, items) -> None:
    hub_w, arrow_w = 0.86, 0.20
    sh = rbox(slide, Inches(x), Inches(y), Inches(hub_w), Inches(h), WHITE, RED, 1.25, 0.14)
    label(sh, [[("rank 级", 9, True, RED)], [("执行图", 9, True, RED)]])
    for p in sh.text_frame.paragraphs:
        p.space_after = 0
    shape(
        slide, MSO_SHAPE.RIGHT_ARROW,
        Inches(x + hub_w + 0.02), Inches(y + h * 0.30),
        Inches(arrow_w - 0.04), Inches(h * 0.40),
        ARROW_GREY, adj=0.5,
    )
    _cards(slide, x + hub_w + arrow_w, y, w - hub_w - arrow_w, h, items)


def _core(slide, x, y, w, h) -> None:
    rbox(slide, Inches(x), Inches(y), Inches(w), Inches(h), RED_SOFT, RED_LINE, 1.0, 0.03)
    pad_x, pad_y = 0.07, 0.05
    ix, iw = x + pad_x, w - 2 * pad_x
    blocks = [
        ("title", 0.20, ("① 多源图捕获与 DAG 归一", STAGE1_HARD)),
        ("cards", 0.40, SOURCES),
        ("arrow", 0.09, None),
        ("flow", 0.40, BIND),
        ("gap", 0.08, None),
        ("title", 0.20, ("② 规模链 / 序列链图变换与域图派生", None)),
        ("lanes", 1.44, None),
        ("merge", 0.09, None),
        ("fork", 0.38, DOMAINS),
        ("gap", 0.08, None),
        ("title", 0.20, ("③ 约束调度与 Trace 重放", STAGE3_HARD)),
        ("flow", 0.40, SCHEDULE),
    ]
    used = sum(b[1] for b in blocks) + pad_y * 2
    slack = h - used
    if slack < -0.005:
        raise RuntimeError(f"DAG engine overflow by {-slack:.3f} in")
    gap_extra = max(slack, 0) / 2
    cursor = y + pad_y
    mid = x + w / 2
    lane_gap = 0.10
    lane_w = (iw - lane_gap) / 2
    for kind, bh, payload in blocks:
        if kind == "gap":
            bh += gap_extra
            _down_arrow(slide, mid, cursor + 0.012, max(bh - 0.024, 0.05))
        elif kind == "title":
            _title(slide, ix, cursor, iw, bh, *payload)
        elif kind == "cards":
            _cards(slide, ix, cursor, iw, bh, payload)
        elif kind == "flow":
            _flow(slide, ix, cursor, iw, bh, payload)
        elif kind == "arrow":
            _down_arrow(slide, mid, cursor + 0.008, bh - 0.016)
        elif kind == "merge":
            # both chains flow into the hub
            for cx in (ix + lane_w / 2, ix + lane_w + lane_gap + lane_w / 2):
                _down_arrow(slide, cx, cursor + 0.008, bh - 0.016)
        elif kind == "lanes":
            _lane(slide, ix, cursor, lane_w, bh, "规模链", SCALE_HARD, SCALE)
            _lane(slide, ix + lane_w + lane_gap, cursor, lane_w, bh, "序列链", SEQ_HARD, SEQ)
        elif kind == "fork":
            _fork(slide, ix, cursor, iw, bh, payload)
        else:
            raise RuntimeError(kind)
        cursor += bh


def _io_row(slide, x, y, w, h, items) -> None:
    _cards(slide, x, y, w, h, items, 10)


def _diagram(slide) -> None:
    tag_x = DIAG_LEFT
    x = tag_x + TAG_W + 0.05
    w = DIAG_RIGHT - x
    core_h = (DIAG_BOTTOM - DIAG_TOP) - INPUT_H - OUTPUT_H - 2 * ARROW
    y = DIAG_TOP
    _tag(slide, tag_x, y, INPUT_H, "输入层", None, SLATE)
    _io_row(slide, x, y, w, INPUT_H, INPUTS)
    y += INPUT_H
    _down_arrow(slide, x + w / 2, y + 0.005, ARROW - 0.01)
    y += ARROW
    _tag(slide, tag_x, y, core_h, "DAG\n引擎", "本页核心", RED)
    _core(slide, x, y, w, core_h)
    y += core_h
    _down_arrow(slide, x + w / 2, y + 0.005, ARROW - 0.01)
    y += ARROW
    _tag(slide, tag_x, y, OUTPUT_H, "输出层", None, SLATE)
    _io_row(slide, x, y, w, OUTPUT_H, OUTPUTS)


def _banner(slide) -> None:
    top, height = 0.78, 1.08
    rbox(slide, Inches(0.22), Inches(top), Inches(12.88), Inches(height), WHITE, LINE, 0.9, 0.05)
    row_h = height / len(BANNER)
    for i, (name, body) in enumerate(BANNER):
        yy = top + i * row_h
        rect(slide, Inches(0.36), Inches(yy + 0.10), Inches(0.07), Inches(row_h - 0.20), RED)
        text(slide, Inches(0.52), Inches(yy), Inches(1.22), Inches(row_h), name,
             12, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        text(slide, Inches(1.78), Inches(yy), Inches(11.10), Inches(row_h), body,
             12, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def _tech_column(slide) -> None:
    left, right = 6.72, 13.08
    top, bottom = DIAG_TOP, DIAG_BOTTOM
    width = right - left
    gap = 0.07
    card_h = (bottom - top - gap * (len(TECHS) - 1)) / len(TECHS)
    text(slide, Inches(left), Inches(top - 0.22), Inches(width), Inches(0.20),
         "关键技术", 12, True, RED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    for i, (kicker, title, points) in enumerate(TECHS):
        yy = top + i * (card_h + gap)
        rbox(slide, Inches(left), Inches(yy), Inches(width), Inches(card_h), WHITE, LINE, 0.9, 0.06)
        rect(slide, Inches(left), Inches(yy + 0.08), Inches(0.07), Inches(card_h - 0.16), RED)
        text(slide, Inches(left + 0.18), Inches(yy + 0.05), Inches(width - 0.30), Inches(0.22),
             kicker, 11, True, RED)
        text(slide, Inches(left + 0.18), Inches(yy + 0.26), Inches(width - 0.30), Inches(0.26),
             title, 13, True, INK)
        slot_top = yy + 0.54
        slot_h = (card_h - 0.62) / 2
        for n, (lead, desc) in enumerate(points):
            parts = desc if isinstance(desc, tuple) else (desc,)
            lines = [[(f"{n + 1}）{lead}", 11, True, INK), (parts[0], 11, False, BODY)]]
            lines += [[(part, 11, False, BODY)] for part in parts[1:]]
            text(
                slide,
                Inches(left + 0.18),
                Inches(slot_top + n * slot_h),
                Inches(width - 0.32),
                Inches(slot_h),
                lines,
                anchor=MSO_ANCHOR.TOP,
            )


def build() -> None:
    prs = new_deck()
    slide = blank(prs)
    rect(slide, 0, 0, prs.slide_width, prs.slide_height, PAPER)
    rect(slide, 0, 0, prs.slide_width, Inches(0.06), RED)
    text(slide, Inches(0.28), Inches(0.12), Inches(12.6), Inches(0.24),
         "设计态 AI 集群仿真器  ·  框架侧仿真能力", 12, True, RED)
    text(slide, Inches(0.28), Inches(0.36), Inches(12.7), Inches(0.38),
         "从多框架导出统一 DAG，并按场景变换为大规模执行图与网络、访存图",
         20, True, INK)
    _banner(slide)
    text(slide, Inches(DIAG_LEFT), Inches(DIAG_TOP - 0.22), Inches(3.4), Inches(0.20),
         "能力结构", 12, True, SLATE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    _diagram(slide)
    _tech_column(slide)
    kit_footer(slide, "框架侧仿真能力  ·  立项讨论稿", 1, 1)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), RED)
    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "v5：框图内只放技术名词，每一步带 1–2 个技术标签，难点直接用红字标出（无边框）；过渡箭头用浅灰；不放句子说明。\n"
        "① 多源图捕获与 DAG 归一：结构配置、源码 AST、框架 IR（FX / Jaxpr / MindIR）、"
        "运行 Trace（Kineto / msprof）并列捕获；再依次做算子实例化（算子分解映射）、"
        "形状精度绑定（符号形状推导）、控制流展开（子图内联）、多框架归一（统一算子 IR）。"
        "难点是动态形状与图断点。\n"
        "② 规模链：设备网格映射（Device Mesh、拓扑亲和放置）→ rank 级展开（SPMD 复制、同构折叠）"
        "→ 并行策略切分（切分传播、集合通信插入）→ 批次并发扩展（微批 1F1B、连续批处理）；"
        "难点是万卡规模下的图膨胀与跨 rank 切分一致性。"
        "序列链：序列并行切分（SP/CP、Ring Attn）→ Prefix 复用（KV 分页、前缀树匹配）"
        "→ 量化算子替换（W8A8/FP8、Q/DQ 插入）→ 融合与重计算（融合 pass、激活重计算）；"
        "难点是变长序列与 KV 显存。"
        "两链汇合为 rank 级执行图，再派生集合通信、点对点流量、内存容量（HBM 生命周期）、数据读取四张域图。\n"
        "③ 约束调度与 Trace 重放：依赖展开（跨 rank 依赖）→ 算子调度（多流、通算重叠）"
        "→ 执行序编排（约束拓扑排序）→ 多场景重放（离散事件重放）；难点是通算重叠与跨 rank 同步。\n"
        "输入直接进入 DAG 引擎，下面是输出层，本页不接仿真核。"
        "配色遵循 AICO-PPT：品牌红 #b5333b 只标 DAG 引擎，结构用灰蓝 #566472。"
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
