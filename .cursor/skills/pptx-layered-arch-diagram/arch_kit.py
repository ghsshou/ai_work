#!/usr/bin/env python3
"""Spec-driven layered architecture diagram for python-pptx (16:9, Chinese fonts).

Usage:
    from arch_kit import new_deck, blank, title_block, footer, draw_layered_arch
    prs = new_deck()
    s = blank(prs)
    title_block(s, "kicker", "title", "subtitle")
    draw_layered_arch(s, SPEC, left=0.36, right=8.86, top=1.2, bottom=6.98)
    prs.save("out.pptx")

SPEC is a list of layers, drawn top to bottom with down arrows between them:
    {"name": "输入层", "sub": None, "tone": "blue", "h": 0.46,
     "items": ["子框1", "子框2"]}                      # simple row of boxes
    {"name": "负载引擎", "sub": "本项目核心", "tone": "core", "h": 3.92,
     "columns": [{"title": "① 方法\\n对象", "items": [...4]}, ...],   # optional
     "columns_h": 1.86,
     "rows": [{"title": "④ ...", "items": [...4], "chevron": True}, ...]}  # optional
    {"name": "仿真内核", "sub": "已具备", "tone": "gray", "items": [...],
     "feedback": True}   # orange up-arrow from this layer into the layer above

Tones: blue / core / gray / green (legacy load-engine palette).
For new charter/立项 slides prefer AICO via aico_tones(): core #b5333b / #fdf0f1.
Also: DIAG_LINE (#5a5a5e), LEAF_FILL (white), BAND_FILLS (soft per-module bands).
"\\n" inside a label forces a centered line break.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "微软雅黑"

BG = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
DIM = RGBColor(0x7A, 0x88, 0x96)
LINE = RGBColor(0xD0, 0xD8, 0xE0)
BLUE = RGBColor(0x1A, 0x5F, 0xA8)
BLUE_BG = RGBColor(0xE3, 0xEE, 0xF8)
CORE = RGBColor(0xC4, 0x5A, 0x10)
CORE_BG = RGBColor(0xFD, 0xF3, 0xEA)
CORE_LT = RGBColor(0xFB, 0xE3, 0xCF)
GREEN = RGBColor(0x3D, 0x86, 0x2C)
GREEN_BG = RGBColor(0xE8, 0xF3, 0xE2)
GRAY_BG = RGBColor(0xE9, 0xEC, 0xF0)

# AICO-PPT charter palette (prefer for new 立项 slides; legacy TONES kept for load-engine example)
AICO_RED = RGBColor(0xB5, 0x33, 0x3B)
AICO_RED_SOFT = RGBColor(0xFD, 0xF0, 0xF1)
AICO_SLATE = RGBColor(0x56, 0x64, 0x72)
DIAG_LINE = RGBColor(0x5A, 0x5A, 0x5E)  # deep gray stroke for modules / leaves / in-out
LEAF_FILL = RGBColor(0xFF, 0xFF, 0xFF)  # white fill for leaf technique + I/O chips
CARD_LINE = RGBColor(0xE7, 0xE7, 0xE7)  # light border for right-side key-tech cards
BAND_FILLS = (
    RGBColor(0xEE, 0xF2, 0xF6),  # cool slate
    RGBColor(0xF7, 0xF0, 0xF1),  # warm rose/cream
    RGBColor(0xEE, 0xF5, 0xF0),  # sage/mint-gray
)

# tone -> (tag color, box fill, box line)
TONES = {
    "blue": (BLUE, BLUE_BG, None),
    "core": (CORE, CORE_BG, CORE_LT),
    "gray": (MUTED, GRAY_BG, LINE),
    "green": (GREEN, GREEN_BG, None),
}


def aico_tones():
    """Return tone map for AICO-PPT charter slides (red core, slate I/O)."""
    return {
        "blue": (AICO_SLATE, LEAF_FILL, DIAG_LINE),
        "core": (AICO_RED, AICO_RED_SOFT, DIAG_LINE),
        "gray": (MUTED, GRAY_BG, DIAG_LINE),
        "green": (AICO_SLATE, LEAF_FILL, DIAG_LINE),
        "aico_core": (AICO_RED, AICO_RED_SOFT, DIAG_LINE),
        "aico_io": (AICO_SLATE, LEAF_FILL, DIAG_LINE),
    }


def band_fill(index: int):
    """Soft fill for the i-th sibling core module band (cycles BAND_FILLS)."""
    return BAND_FILLS[index % len(BAND_FILLS)]


def leaf_chip(slide, l, t, w, h, content, size=8.8, bold=False, color=INK,
              fill=None, line=None, line_pt=0.75):
    """Rounded leaf / I/O chip: white fill + deep-gray stroke by default."""
    fill = LEAF_FILL if fill is None else fill
    line = DIAG_LINE if line is None else line
    sh = rbox(slide, l, t, w, h, fill, line, line_pt, 0.12)
    label(sh, content, size, bold, color)
    return sh


W = Inches(13.333)
H = Inches(7.5)
TAG_W = Inches(0.78)
ARROW_GAP = Inches(0.2)


def _ea_font(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", FONT)


def _style(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    _ea_font(run)


def _frame(tf, anchor, lr=0.05, tb=0.03):
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(lr)
    tf.margin_top = tf.margin_bottom = Inches(tb)
    tf._txBody.bodyPr.set("anchor", {
        MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b",
    }[anchor])


def _fill_lines(tf, content, size, bold, color, align):
    """content: str (\\n = new centered paragraph) or list of lines of (text, size, bold, color) segments."""
    lines = content if isinstance(content, list) else content.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        segs = line if isinstance(line, list) else [(line, size, bold, color)]
        for seg in segs:
            run = p.add_run()
            run.text = seg[0]
            _style(run, *seg[1:])


def text(slide, l, t, w, h, content, size=11, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    _frame(box.text_frame, anchor)
    _fill_lines(box.text_frame, content, size, bold, color, align)
    return box


def shape(slide, kind, l, t, w, h, fill, line=None, line_pt=0.9, adj=None):
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_pt)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.shadow.inherit = False
    return sh


def rect(slide, l, t, w, h, fill, line=None):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line)


def rbox(slide, l, t, w, h, fill, line=None, line_pt=0.9, adj=0.12):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line, line_pt, adj)


def label(sh, content, size=8.8, bold=False, color=INK):
    _frame(sh.text_frame, MSO_ANCHOR.MIDDLE, 0.03, 0.01)
    _fill_lines(sh.text_frame, content, size, bold, color, PP_ALIGN.CENTER)


def grid(slide, l, t, w, h, labels, cols, fill, line, size=8.8, gap=0.06):
    rows = -(-len(labels) // cols)
    g = Inches(gap)
    cw = (w - (cols - 1) * g) / cols
    rh = (h - (rows - 1) * g) / rows
    for i, lab in enumerate(labels):
        r, c = divmod(i, cols)
        label(rbox(slide, l + c * (cw + g), t + r * (rh + g), cw, rh, fill, line, 0.75), lab, size)


def layer_tag(slide, l, t, h, name, sub, color):
    rect(slide, l, t, TAG_W, h, color)
    lines = [[(name, 10.5, True, WHITE)]] + ([[(sub, 8, False, WHITE)]] if sub else [])
    text(slide, l, t, TAG_W, h, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x, t, up=False, color=DIM):
    kind = MSO_SHAPE.UP_ARROW if up else MSO_SHAPE.DOWN_ARROW
    shape(slide, kind, x - Inches(0.16), t, Inches(0.32), Inches(0.14), color, adj=0.5)


def module_column(slide, l, t, w, h, title, items, accent, fill, line, title_h=Inches(0.44),
                  cols=None, leaf_fill=None, leaf_line=None):
    """cols=None -> 1 column (legacy); cols=4 with 8 items draws a 2x4 density grid."""
    stroke = leaf_line if leaf_line is not None else (line if line is not None else accent)
    rbox(slide, l, t, w, h, WHITE, stroke, 1.0, 0.04)
    text(slide, l + Inches(0.05), t + Inches(0.02), w - Inches(0.1), title_h, title,
         10, True, accent, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pad = Inches(0.07)
    n_cols = cols if cols is not None else 1
    chip_fill = leaf_fill if leaf_fill is not None else fill
    chip_line = leaf_line if leaf_line is not None else line
    grid(slide, l + pad, t + title_h + Inches(0.02), w - 2 * pad, h - title_h - Inches(0.09),
         items, n_cols, chip_fill, chip_line)


def module_row(slide, l, t, w, h, title, items, accent, fill, line,
               title_w=Inches(1.62), chevron=False):
    rbox(slide, l, t, w, h, WHITE, accent, 1.0, 0.08)
    text(slide, l + Inches(0.06), t, title_w - Inches(0.06), h, title,
         10, True, accent, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pad = Inches(0.07)
    x0 = l + title_w + Inches(0.04)
    iw = w - title_w - Inches(0.04) - pad
    if not chevron:
        grid(slide, x0, t + pad, iw, h - 2 * pad, items, len(items), fill, line)
        return
    cw = iw / len(items)
    for i, lab in enumerate(items):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = shape(slide, kind, x0 + i * cw, t + pad, cw + Inches(0.05), h - 2 * pad,
                   fill, line, 0.75, adj=0.22)
        label(sh, lab)


def _draw_container(slide, x0, t, w, h, layer, accent, fill, line):
    # Outer envelope stroke: deep gray when layer requests charter style.
    envelope_line = layer.get("stroke", accent)
    rbox(slide, x0, t, w, h, fill, envelope_line, 1.75, 0.02)
    pad, gap = Inches(0.08), Inches(0.07)
    ix, iw = x0 + pad, w - 2 * pad
    y = t + pad
    cols = layer.get("columns") or []
    rows = layer.get("rows") or []
    leaf_fill = layer.get("leaf_fill", fill)
    leaf_line = layer.get("leaf_line", line)
    if cols:
        ch = Inches(layer.get("columns_h", 1.86)) if rows else h - 2 * pad
        mw = (iw - (len(cols) - 1) * gap) / len(cols)
        for i, m in enumerate(cols):
            mod_fill = m.get("band_fill", band_fill(i) if layer.get("banded") else fill)
            # Optional per-module tinted band behind the column card.
            if layer.get("banded") or "band_fill" in m:
                rbox(slide, ix + i * (mw + gap), y, mw, ch, mod_fill, leaf_line or DIAG_LINE, 1.0, 0.04)
            module_column(
                slide, ix + i * (mw + gap), y, mw, ch, m["title"], m["items"],
                accent, leaf_fill, leaf_line,
                cols=m.get("cols", layer.get("item_cols")),
                leaf_fill=m.get("leaf_fill", leaf_fill),
                leaf_line=m.get("leaf_line", leaf_line),
            )
        y += ch + gap
    if rows:
        rh = (t + h - pad - y - (len(rows) - 1) * gap) / len(rows)
        for m in rows:
            module_row(slide, ix, y, iw, rh, m["title"], m["items"], accent,
                       leaf_fill, leaf_line, chevron=m.get("chevron", False))
            y += rh + gap


def draw_layered_arch(slide, spec, left=0.36, right=12.97, top=1.2, bottom=None):
    """Draw layers top to bottom inside [left, right] inches. Returns bottom y (EMU)."""
    tag_l = Inches(left)
    x0 = tag_l + TAG_W + Inches(0.06)
    x1 = Inches(right)
    w = x1 - x0
    y = Inches(top)
    if bottom is not None:
        fixed = sum(Inches(L.get("h", 0.46)) for L in spec) + ARROW_GAP * (len(spec) - 1)
        if fixed > Inches(bottom) - y:
            raise ValueError(f"layers need {fixed / 914400:.2f}in, only {bottom - top:.2f}in available")
    prev = None
    for L in spec:
        accent, fill, line = TONES[L.get("tone", "blue")]
        h = Inches(L.get("h", 0.46))
        if prev is not None:
            arrow(slide, x0 + w / 2, y - ARROW_GAP + Inches(0.03))
            if L.get("feedback"):
                arrow(slide, x1 - Inches(0.8), y - ARROW_GAP + Inches(0.03), up=True, color=CORE)
        layer_tag(slide, tag_l, y, h, L["name"], L.get("sub"), accent)
        if L.get("columns") or L.get("rows"):
            _draw_container(slide, x0, y, w, h, L, accent, fill, line)
        else:
            items = L["items"]
            grid(slide, x0, y, w, h, items, L.get("cols", len(items)), fill, line,
                 L.get("size", 8.8), 0.07)
        prev = L
        y += h + ARROW_GAP
    return y - ARROW_GAP


def new_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, BG)
    return s


def title_block(slide, kicker, title, subtitle=None, width=12.6):
    rect(slide, 0, 0, W, Inches(0.07), BLUE)
    text(slide, Inches(0.36), Inches(0.13), Inches(width), Inches(0.26), kicker, 10.5, True, BLUE)
    text(slide, Inches(0.36), Inches(0.36), Inches(width), Inches(0.46), title, 21, True, INK)
    if subtitle:
        text(slide, Inches(0.36), Inches(0.82), Inches(width), Inches(0.32), subtitle, 11.5, False, MUTED)


def footer(slide, note, page=None, total=None):
    rect(slide, 0, Inches(7.26), W, Inches(0.24), GRAY_BG)
    rect(slide, 0, Inches(7.26), Inches(0.16), Inches(0.24), BLUE)
    text(slide, Inches(0.36), Inches(7.26), Inches(11), Inches(0.24), note, 8.5, False, DIM,
         anchor=MSO_ANCHOR.MIDDLE)
    if page:
        text(slide, Inches(11.3), Inches(7.26), Inches(1.7), Inches(0.24), f"{page:02d}  /  {total:02d}",
             8.5, False, MUTED, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
